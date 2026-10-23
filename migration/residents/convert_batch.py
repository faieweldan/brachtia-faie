"""
Part 1 of the resident migration: turn 'Migration Batch 1' into the Sheet3 upload.

    Master workbook -> Migration Batch 1 -> Sheet3 columns -> a new prepared workbook

It never touches the database and never changes either input file. What each
column means lives in src/lib/migration/sheet3-contract.json - this script only
follows it. The console prints counts, never a resident's details.

Run:
    python3 migration/residents/convert_batch.py
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import unicodedata
import warnings
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import urlparse

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# a broken date in a column the converter never reads (Google Form Q366) - not worth a scary line
warnings.filterwarnings("ignore", message=".*is marked as a date but the serial value.*")

ROOT = Path(__file__).resolve().parents[2]
CONTRACT = json.loads((ROOT / "src/lib/migration/sheet3-contract.json").read_text())

DEFAULT_SOURCE = ROOT / "references/Brachtia Homes_Master_2026 (1).xlsx"
DEFAULT_LABELS = ROOT / "references/Data Labelling.xlsx"
DEFAULT_APPLICATIONS = ROOT / "references/DB Restructure - BMR.xlsx"
DEFAULT_OUT = ROOT / "migration/residents/output/Data Labelling - Batch 1 Prepared.xlsx"
DEFAULT_REPORT = ROOT / "migration/residents/output/Batch 1 Conversion Report.xlsx"
# answers from the team, applied on top of Batch 1 until the Google Sheet is fixed
DEFAULT_CORRECTIONS = ROOT / "references/batch1-corrections.json"

# How serious a finding is. The upload's own Ready / Warning / Blocked comes from these.
BLOCKED, REVIEW, WARNING = "Blocked", "Needs review", "Warning"

# Country calling codes. They are prefix-free by design (no code is the start of
# another), so reading digits left to right finds exactly one.
CALLING_CODES = set(
    "1 7 20 27 30 31 32 33 34 36 39 40 41 43 44 45 46 47 48 49 51 52 53 54 55 56 57 58 "
    "60 61 62 63 64 65 66 81 82 84 86 90 91 92 93 94 95 98 211 212 213 216 218 220 221 222 "
    "223 224 225 226 227 228 229 230 231 232 233 234 235 236 237 238 239 240 241 242 243 "
    "244 245 246 248 249 250 251 252 253 254 255 256 257 258 260 261 262 263 264 265 266 "
    "267 268 269 290 291 297 298 299 350 351 352 353 354 355 356 357 358 359 370 371 372 "
    "373 374 375 376 377 378 380 381 382 383 385 386 387 389 420 421 423 500 501 502 503 "
    "504 505 506 507 508 509 590 591 592 593 594 595 596 597 598 599 670 672 673 674 675 "
    "676 677 678 679 680 681 682 683 685 686 687 688 689 690 691 692 850 852 853 855 856 "
    "880 886 960 961 962 963 964 965 966 967 968 970 971 972 973 974 975 976 977 992 993 "
    "994 995 996 998".split()
)

# Batch 1's spellings -> the value the system stores (the "Value map" tab of Data
# Labelling, 12 Oct 2026). Whatever comes out is then checked against "Allowed
# values" - a value not on that list is left blank and flagged, never invented.
FREQUENCY = {"MONTHLY": "monthly", "BI-MONTHLY": "bimonthly", "QUARTERLY": "quarterly",
             "HALF-YEARLY": "semiannual", "SEMI-ANNUALLY": "semiannual", "ANNUALLY": "annual", "FULLY": "full"}
GENDER = {"F": "Female", "M": "Male", "FEMALE": "Female", "MALE": "Male"}
UNIVERSITY = {"UOC": "UoC", "UNIVERSITY OF CYBERJAYA": "UoC", "CTU": "CU", "CITYU": "CU",
              "CITY UNIVERSITY": "CU", "MULTIMEDIA UNIVERSITY": "MMU", "HERIOT-WATT": "HWUM"}
RESIDENCE = {"ARC": "The Arc, Cyberjaya", "THE ARC": "The Arc, Cyberjaya",
             "SOLSTICE": "Solstice Residence, Cyberjaya", "SOLSTICE RESIDENCE": "Solstice Residence, Cyberjaya"}
PAID_BY = {"SELF": "Self", "SPONSOR": "Sponsor", "PARENT": "Parent", "GUARDIAN": "Guardian"}

# Sheet3 column -> its list in the "Allowed values" tab
DROPDOWNS = {"Gender": "Gender", "Nationality": "Nationality", "Resident type": "Resident type",
             "University": "University (application form)", "Marital status": "Marital status",
             "Race": "Race", "Religion": "Religion", "Residence": "Residence",
             "Payment frequency": "Payment frequency", "Paid by": "Paid by", "Room": "Room", "Bed": "Bed",
             "Referral Source": "Referral source", "Medical Condition": "Medical condition"}


def key(header) -> str:
    """'Resident name', 'Resident Name' and 'Unit no. ' are the same column."""
    return " ".join(str(header or "").lower().split())


def allowed_values(labels_path: Path) -> dict[str, set[str]]:
    """field -> stored values, from the 'Allowed values' tab. Empty when the tab is not there."""
    wb = openpyxl.load_workbook(labels_path, read_only=True)
    if "Allowed values" not in wb.sheetnames:
        wb.close()
        return {}
    out: dict[str, set[str]] = {}
    for field, value, *_ in wb["Allowed values"].iter_rows(min_row=4, values_only=True):
        if field and value not in (None, ""):
            out.setdefault(str(field), set()).add(str(value))
    wb.close()
    return out
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def fingerprint(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(v) -> str:
    """A cell as text, without the '.0' Excel puts on a number typed into an ID column."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    # invisible characters pasted in with a name (a zero-width space) are not kept
    return re.sub(r"[\u200b-\u200d\u2060\ufeff]", "", str(v)).strip()


# ---------------------------------------------------------------- headers

def source_columns(ws) -> dict[tuple[str, int], int]:
    """(header, occurrence) -> column number. Two 'Status' columns become ('Status', 1) and ('Status', 2)."""
    seen: dict[str, int] = {}
    out = {}
    for cell in ws[1]:
        name = text(cell.value)
        if not name:
            continue
        seen[name] = seen.get(name, 0) + 1
        out[(name, seen[name])] = cell.column
    return out


def sheet3_headers(labels_path: Path) -> list[str]:
    """Only the header row of Sheet3 is read - nothing below it."""
    wb = openpyxl.load_workbook(labels_path, read_only=True)
    ws = wb[CONTRACT["targetSheet"]]
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    wb.close()
    return [h for h in row if h is not None]


def upload_headers(sheet3: list[str]) -> list[str]:
    """Sheet3's order, followed by trace-only migration columns."""
    known = {key(c["sheet3"]) for c in CONTRACT["columns"]}
    unknown = [h for h in sheet3 if key(h) not in known]
    if unknown:
        raise SystemExit(f"Sheet3 has columns the contract does not describe: {unknown}")
    have = {key(h) for h in sheet3}
    missing = [c["sheet3"] for c in CONTRACT["columns"] if not c.get("addedToSheet3") and key(c["sheet3"]) not in have]
    if missing:
        raise SystemExit(f"The contract expects Sheet3 columns that are not there: {missing}")
    added = [c["sheet3"] for c in CONTRACT["columns"] if c.get("addedToSheet3") and key(c["sheet3"]) not in have]
    traces = [c["sheet3"] for c in CONTRACT["migrationOnly"] if key(c["sheet3"]) not in have]
    return list(sheet3) + added + traces


# ---------------------------------------------------------------- application form

APPLICATION_IDENTITY = {
    "document_id": "Passport / NRIC Number",
    "student_id": "Student ID",
    "email": "Email address",
    "phone": "Contact Number (Mobile)",
    "name": "Full Name",
}


def normalized_email(value) -> str:
    return text(value).lower()


def normalized_name(value) -> str:
    """Case, accents and punctuation do not stop an otherwise exact name match."""
    raw = unicodedata.normalize("NFKD", text(value).casefold())
    ascii_name = "".join(c for c in raw if not unicodedata.combining(c))
    return " ".join(re.findall(r"[\w]+", ascii_name, flags=re.UNICODE))


def normalized_phone(value) -> str:
    """Comparable international digits; a local Malaysian 01x becomes 60 1x."""
    digits = re.sub(r"\D", "", text(value))
    if digits.startswith("00"):
        digits = digits[2:]
    if digits.startswith("01"):
        digits = "60" + digits[1:]
    return digits.lstrip("0")


def normalized_identifier(value) -> str:
    """Exact ID comparison after removing case and presentation punctuation."""
    compact = re.sub(r"[^0-9A-Z]", "", unicodedata.normalize("NFKC", text(value)).upper())
    if compact in {"NA", "NONE", "NIL", "NOTAPPLICABLE", "NOTAVAILABLE", "TBC", "TBD"}:
        return ""
    return compact


def application_rows(path: Path) -> list[dict]:
    """Read the Google Form once. Workbook instructions are data, never executable."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = CONTRACT.get("applicationSheet", "gform responses")
    if sheet not in wb.sheetnames:
        wb.close()
        raise SystemExit(f"Application workbook has no '{sheet}' sheet")
    ws = wb[sheet]
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    columns = {key(value): i for i, value in enumerate(first) if text(value)}
    required = set(APPLICATION_IDENTITY.values())
    for column in CONTRACT["columns"]:
        sources = column.get("applicationSource") or []
        required.update([sources] if isinstance(sources, str) else sources)
    absent = sorted(name for name in required if key(name) not in columns)
    if absent:
        wb.close()
        raise SystemExit(f"Application form is missing expected columns: {absent}")

    rows = []
    for row_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value not in (None, "") for value in values):
            continue
        get = lambda name: values[columns[key(name)]]
        stamp = get("Timestamp") if key("Timestamp") in columns else None
        rows.append({
            "_row": row_no,
            "_timestamp": stamp if isinstance(stamp, datetime) else datetime.min,
            "_document_id": normalized_identifier(get(APPLICATION_IDENTITY["document_id"])),
            "_student_id": normalized_identifier(get(APPLICATION_IDENTITY["student_id"])),
            "_email": normalized_email(get(APPLICATION_IDENTITY["email"])),
            "_phone": normalized_phone(get(APPLICATION_IDENTITY["phone"])),
            "_name": normalized_name(get(APPLICATION_IDENTITY["name"])),
            "_values": {name: values[index] for name, index in columns.items()},
        })
    wb.close()
    return rows


def application_indexes(rows: list[dict]) -> dict[str, dict[str, list[dict]]]:
    indexes = {field: defaultdict(list) for field in APPLICATION_IDENTITY}
    for row in rows:
        for field in indexes:
            value = row.get(f"_{field}", "")
            if value:
                indexes[field][value].append(row)
    return indexes


def match_application(resident: dict, indexes: dict[str, dict[str, list[dict]]]):
    """Return (row, evidence, finding). One reliable identity may match; conflicts stay visible."""
    identity = {
        "document_id": normalized_identifier(
            resident.get("NRIC number") or resident.get("Passport number") or resident.get("NRIC / Passport")
        ),
        # Batch 1's StudentID is the legacy QuickBooks ID, not the university
        # student number. Only compare a separate university Student ID when a
        # source provides one; never conflate the two identifiers.
        "student_id": normalized_identifier(resident.get("Student ID") or resident.get("University Student ID")),
        "email": normalized_email(resident.get("Email")),
        "phone": normalized_phone(f"{resident.get('Country Code', '')}{resident.get('Mobile Number', '')}"),
        "name": normalized_name(resident.get("Resident Name")),
    }
    support: dict[int, set[str]] = defaultdict(set)
    candidates: dict[int, dict] = {}
    for field, value in identity.items():
        if not value:
            continue
        for row in indexes[field].get(value, []):
            support[row["_row"]].add(field)
            candidates[row["_row"]] = row
    if not support:
        # not a stopper: after import, admin sends the Application Form Link and the resident fills it in (Dani, 16 Oct 2026)
        return None, "No match", (
            WARNING,
            "No application form matches the NRIC/passport, university Student ID, email, phone or normalized name"
            " - after import, send them the Application Form Link",
        )

    best = max(len(evidence) for evidence in support.values())
    winners = [candidates[row_no] for row_no, evidence in support.items() if len(evidence) == best]
    labels = {
        "document_id": "NRIC/passport",
        "student_id": "university Student ID",
        "email": "email",
        "phone": "phone",
        "name": "normalized name",
    }
    evidence = sorted(support[winners[0]["_row"]])
    method = " + ".join(labels[field] for field in evidence)

    if len(winners) > 1:
        if best < 2:
            rows = ", ".join(str(row["_row"]) for row in winners)
            return None, f"Ambiguous: {method}", (BLOCKED, f"Application form match is ambiguous between rows {rows}")
        chosen = max(winners, key=lambda row: (row["_timestamp"], row["_row"]))
        # sent the form more than once: the newest answers are theirs (Dani, 13 Oct 2026)
        return chosen, f"{method}; newest of rows {', '.join(str(r['_row']) for r in winners)}", None

    chosen = winners[0]
    others = sorted(row_no for row_no in support if row_no != chosen["_row"])
    if others and best >= 2:
        # two or more fields agree: that row is theirs; a lone stray match (a shared
        # phone) is noted in the matches tab, not raised - staff fix it on the profile
        return chosen, f"{method}; also weakly matched row(s) {', '.join(map(str, others))}", None
    if others:
        return None, f"Ambiguous: {method}", (
            BLOCKED,
            f"Different identity fields point to application rows {chosen['_row']}, {', '.join(map(str, others))}",
        )
    if best == 1 and evidence == ["name"]:
        return chosen, method, (WARNING, "Application form matched only by normalized name - confirm before import")
    return chosen, method, None


def application_value(row: dict, header: str):
    return row["_values"].get(key(header))


def normal_race(value) -> tuple[str, str]:
    raw = text(value)
    k = key(raw)
    if not k.strip(".-/ "):
        return "", ""
    mapped = {
        "malay": "Malay", "melayu": "Malay", "chinese": "Chinese", "indian": "Indian",
        # Sponsor forms (20 Oct 2026): East Malaysian groups are this option
        "iban": "Bumiputera (Sabah/Sarawak)", "kadazan": "Bumiputera (Sabah/Sarawak)",
        "bajau": "Bumiputera (Sabah/Sarawak)",
        "bumiputera sabah, brunei": "Bumiputera (Sabah/Sarawak)",
        "bumiputera (sabah/sarawak)": "Bumiputera (Sabah/Sarawak)",
        "bamar": "Other Asian", "burmese": "Other Asian",
        "black african": "African", "african": "African",
        "middle eastern": "Middle Eastern", "arab": "Middle Eastern", "arabian": "Middle Eastern",
        "cina": "Chinese", "indonesian": "Other Asian", "asian": "Other Asian", "malagasy": "African", "n\\a": "Prefer not to say", "european": "European", "white": "European", "caucasian": "European",
        "other": "Other", "n/a": "Prefer not to say", "n": "Prefer not to say",
        "prefer not to say": "Prefer not to say",
    }.get(k)
    if not mapped and ("sabah" in k or "sarawak" in k):
        mapped = "Bumiputera (Sabah/Sarawak)"
    return (mapped, "") if mapped else ("Other", f"Race '{raw}' was grouped as Other")


def normal_religion(value) -> tuple[str, str]:
    raw = text(value)
    k = key(raw).strip(". ")
    mapped = {
        "islam": "Islam", "muslim": "Islam", "islam/musllim": "Islam",
        "christian": "Christianity", "christian chatolic": "Christianity", "christian catholic": "Christianity", "christianity": "Christianity", "kristian": "Christianity",
        "buddha": "Buddhism", "buddhist": "Buddhism", "buddhism": "Buddhism",
        "hindu": "Hinduism", "hinduism": "Hinduism",
        "sikh": "Sikhism", "sikhism": "Sikhism", "tao": "Taoism", "taoism": "Taoism",
        "none": "No Religion / None", "no": "No Religion / None", "atheist": "No Religion / None", "no religion": "No Religion / None", "agnostic": "No Religion / None", "non practicing": "No Religion / None",
        "n/a": "Prefer not to say", "prefer not to say": "Prefer not to say", "other": "Other",
    }.get(k)
    return (mapped, "") if mapped else ("Other", f"Religion '{raw}' was grouped as Other")


def normal_marital(value) -> tuple[str, str]:
    """The form's answer -> the system's marital status; anything else is left blank, never invented."""
    raw = text(value)
    k = key(raw).strip(". ")
    if not k:
        return "", ""
    mapped = {
        "single": "Single", "bujang": "Single", "unmarried": "Single", "not married": "Single",
        "married": "Married", "berkahwin": "Married",
        "divorced": "Divorced", "bercerai": "Divorced",
        "widowed": "Widowed", "widow": "Widowed", "widower": "Widowed",
    }.get(k)
    return (mapped, "") if mapped else ("", f"Marital status '{raw}' is not on the list - left blank")


def normal_referral(value) -> tuple[str, str]:
    """How they heard about us -> the application form's list (HEARD_ABOUT); unknown is Other."""
    raw = text(value)
    k = key(raw)
    if not k.strip(".-/ "):
        return "", ""
    for words, option in (
        (("universit", "uni ", "campus", "college", "lecturer", "faculty",
          "mmu", "hwum", "heriot", "uoc", "cyberjaya uni", "city u"), "University"),
        (("google", "search", "website", "internet", "online"), "Google Search"),
        (("friend", "senior", "family", "sibling", "brother", "sister", "cousin", "relative", "parent", "word of mouth"), "Friend"),
        (("agent", "agency"), "Agent"),
        (("social", "instagram", "facebook", "tiktok", "twitter", "xiaohongshu", "whatsapp", "telegram"), "Social Media"),
        (("other",), "Other"),
    ):
        if any(w in f"{k} " for w in words):
            return option, ""
    return "Other", f"Referral source '{raw}' was grouped as Other"


def normal_medical(value) -> tuple[str, str]:
    """Y/N only (Sheet2): stored as the system's yes / no."""
    raw = text(value)
    k = key(raw).strip(". ")
    if not k:
        return "", ""
    if k in ("yes", "y", "ya"):
        return "yes", ""
    if k in ("no", "n", "none", "nil", "tiada", "tidak", "n/a", "na"):
        return "no", ""
    return "", f"Medical condition '{raw}' is not Yes or No - left blank"


def enrich_from_application(out_row: dict, application: dict, base: Path, flag):
    """Copy only the Sheet2 fields the user named; leave every other form answer out."""
    out_row["Application Source Row"] = application["_row"]
    out_row["Referral Source"], referral_note = normal_referral(application_value(application, "How did you hear about us"))
    out_row["Marital status"], marital_note = normal_marital(application_value(application, "Marital Status"))
    if referral_note:
        flag(WARNING, "Referral Source", referral_note)
    if marital_note:
        flag(WARNING, "Marital status", marital_note)

    old_race = text(application_value(application, "Race"))
    new_race = text(application_value(application, "Race 2"))
    race_raw = old_race if old_race and key(new_race) in ("", "n/a", "n") else new_race or old_race
    out_row["Race"], race_note = normal_race(race_raw)
    out_row["Religion"], religion_note = normal_religion(application_value(application, "Religion"))
    if race_note:
        flag(WARNING, "Race", race_note)
    if religion_note:
        flag(WARNING, "Religion", religion_note)

    for target, source in (
        ("Address", "Residential Address"),
        ("Postal Code", "Postal Code"),
        ("State / Region / Province", "State/ Region/ Province"),
        ("Country", "Country"),
        ("Medical Condition - Remarks", 'If you answered "Yes" to the question above, please specify'),
    ):
        out_row[target] = text(application_value(application, source))
    out_row["Medical Condition"], medical_note = normal_medical(
        application_value(application, "Do you have any medical conditions we should be aware of?"))
    if medical_note:
        flag(WARNING, "Medical Condition", medical_note)
    if out_row["Medical Condition"] == "no":
        out_row["Medical Condition - Remarks"] = ""

    for target, source, label in (
        ("Photograph", "Please upload a colour passport size photograph of yourself", "Photograph"),
        ("NRIC / Passport Copy", "Please upload a colour copy of your NRIC (for Malaysians only) or Passport Information Page", "NRIC / Passport copy"),
        ("Offer Letter", "Please upload your university offer letter", "Offer letter"),
    ):
        value, finding = check_document(text(application_value(application, source)), None, base, label, True)
        out_row[target] = value
        if finding:
            flag(finding[0], target, finding[1])


def enrich_from_answers(out_row: dict, answers: dict, flag):
    """The same fields as the Google Form, from the team's answers - through the
    same mappings, so they land on the same dropdown values. No documents."""
    out_row["Intake"] = text(answers.get("Intake"))
    for header, fn in (("Referral Source", normal_referral), ("Marital status", normal_marital),
                       ("Race", normal_race), ("Religion", normal_religion), ("Medical Condition", normal_medical)):
        out_row[header], note = fn(answers.get(header))
        if note:
            flag(WARNING, header, note)
    for header in ("Address", "Postal Code", "State / Region / Province", "Country", "Medical Condition - Remarks"):
        out_row[header] = text(answers.get(header))
    if out_row["Medical Condition"] == "no":
        out_row["Medical Condition - Remarks"] = ""
    # document links, when the team has them; any not given is uploaded on the profile later
    missing = []
    for header in ("Photograph", "NRIC / Passport Copy", "Offer Letter"):
        out_row[header] = text(answers.get(header))
        if not out_row[header]:
            missing.append(header)
    flag(WARNING, "Application Form", "No Google Form - details from the team's answer"
         + (f"; upload on the profile later: {', '.join(missing)}" if missing else ""))


# ---------------------------------------------------------------- field rules

def split_phone(raw: str, nationality: str = ""):
    """-> (country code like '+60', digits, finding or None)."""
    if not raw:
        return "", "", (BLOCKED, "Mobile number is blank")
    # Excel needs a ' to keep a number as typed - it is not part of the number
    compact = re.sub(r"[\s\-().'\u2019]", "", raw)
    if compact.startswith("+") or compact.startswith("00"):
        digits = compact.lstrip("+")
        if digits.startswith("00"):
            digits = digits[2:]
        if not digits.isdigit():
            return "", "", (REVIEW, "Mobile number has characters that are not digits")
        for n in (1, 2, 3):
            if digits[:n] in CALLING_CODES:
                rest = digits[n:].lstrip("0")
                if not 6 <= len(rest) <= 12:
                    return f"+{digits[:n]}", rest, (REVIEW, "Mobile number length looks wrong")
                return f"+{digits[:n]}", rest, None
        return "", "", (REVIEW, "Mobile number has a country code that is not recognised")
    if compact.isdigit() and compact.startswith("01"):
        # 01x is how a Malaysian mobile is written at home
        return "+60", compact[1:], (WARNING, "No country code - taken as Malaysian (+60) from the 01x format")
    if compact.isdigit() and compact.startswith("601") and 11 <= len(compact) <= 12:
        # +601x... with its plus lost: a Malaysian mobile whoever holds it (a
        # foreign student on a Malaysian line). 601 + 8-9 digits cannot be an
        # Indian 10-digit number, so this is safe for every nationality.
        return "+60", compact[2:], (WARNING, "Number started with 601 - a Malaysian mobile, split into +60 and the number")
    if nationality == "MYS" and compact.isdigit() and compact.startswith("60") and 10 <= len(compact) <= 12:
        # Excel often turns +6011... into the number 6011..., losing the plus.
        return "+60", compact[2:].lstrip("0"), (WARNING, "Malaysian number started with 60 - split into +60 and the mobile number")
    if nationality == "MYS" and compact.isdigit() and compact.startswith("1") and 9 <= len(compact) <= 10:
        # A local 01x number stored as a number loses its leading zero.
        return "+60", compact, (WARNING, "Malaysian local number lost its leading zero - taken as +60")
    dial = NATIONAL_DIAL.get(nationality, "")
    if dial and dial != "60" and compact.isdigit() and compact.startswith(dial) and 6 <= len(compact) - len(dial) <= 12:
        # the + was lost but the number starts with their own country's code
        return f"+{dial}", compact[len(dial):].lstrip("0"), (
            WARNING, f"No + - the number starts with {nationality}'s country code, split into +{dial} and the number")
    return "", "", (REVIEW, "Mobile number has no country code and is not a Malaysian 01x number")


def split_id(nationality: str, raw_value):
    """-> (nric, passport, finding or None). Malaysians carry an NRIC, everyone else a passport."""
    raw = text(raw_value)
    if not re.fullmatch(r"[A-Z]{3}", nationality):
        return "", "", (REVIEW, "Nationality missing or unclear, so the ID type cannot be told")
    if not raw:
        return "", "", (REVIEW, "ID number is blank")
    if nationality == "MYS":
        digits = raw.replace("-", "")
        if isinstance(raw_value, (int, float)) and len(digits) in (10, 11):
            # Excel dropped the zeros of a 2000s birth year (Dani, 12 Oct 2026):
            # put them back, but only when the first six digits are then a real date
            padded = digits.zfill(12)
            yy, mm, dd = int(padded[:2]), int(padded[2:4]), int(padded[4:6])
            if 1 <= mm <= 12 and 1 <= dd <= 31:
                return (f"{padded[:6]}-{padded[6:8]}-{padded[8:]}", "",
                        (WARNING, "NRIC lost its leading zero in Excel - the zero was put back"))
        if isinstance(raw_value, (int, float)) and len(digits) < 12:
            return "", "", (REVIEW, "NRIC was stored as a number and has lost its leading zero - retype it as text")
        if not re.fullmatch(r"\d{12}", digits):
            return "", "", (REVIEW, "Malaysian resident, but the ID is not a 12-digit NRIC")
        return f"{digits[:6]}-{digits[6:8]}-{digits[8:]}", "", None
    return "", raw, None


def legacy_agreement_no(start, unit: str, qb: str) -> str:
    """An existing resident's TA number, e.g. TA2025_A-18-05_00724; blank when a part is missing.
    New residents from a booking get the system's own TA-<Resident ID>-01 instead."""
    if not start or not unit or not qb:
        return ""
    return f"TA{start.year}_{unit}_{qb}"


def stay_duration(start: date, end: date) -> str:
    """What the sheet's DATEDIF(start, end+1) shows, worked out here rather than by running the formula."""
    e = end + timedelta(days=1)
    months = (e.year - start.year) * 12 + e.month - start.month
    if e.day < start.day:
        months -= 1
    y, m = divmod(start.month - 1 + months, 12)
    anchor_year, anchor_month = start.year + y, m + 1
    last = (date(anchor_year + anchor_month // 12, anchor_month % 12 + 1, 1) - timedelta(days=1)).day
    anchor = date(anchor_year, anchor_month, min(start.day, last))
    return f"{months} Months {(e - anchor).days} Days"


def add_months(value: date, months: int) -> date:
    """Excel EDATE behaviour: keep the day, or use the target month's final day."""
    y, m = divmod(value.month - 1 + months, 12)
    year, month = value.year + y, m + 1
    last = (date(year + month // 12, month % 12 + 1, 1) - timedelta(days=1)).day
    return date(year, month, min(value.day, last))


def check_document(raw: str, hyperlink: str | None, base: Path, label: str, required: bool):
    """-> (value to write, finding or None). The link behind the cell wins over its display text."""
    value = (hyperlink or raw or "").strip()
    if not value:
        return "", ((BLOCKED if required else WARNING), f"{label} is missing")
    if re.match(r"^[a-z]+://", value, re.I):
        u = urlparse(value)
        if u.scheme not in ("http", "https") or not u.netloc:
            return value, (BLOCKED, f"{label} is not a valid web link")
        return value, None
    path = Path(value) if Path(value).is_absolute() else base / value
    if not path.exists():
        return value, (BLOCKED, f"{label} points to a local file that does not exist")
    return value, None


RANK = {"Ready": 0, WARNING: 1, REVIEW: 2, BLOCKED: 3}

# A finding's field -> the Migration Batch 1 column to fix it in, so the report
# can say "N · IDNumber" rather than make someone hunt across 37 columns
FIELD_SOURCE = {
    "QuickBooks ID": "StudentID", "Resident Name": "StudentName", "Gender": "Gender",
    "Mobile Number": "MobileNumber", "Email": "Email", "NRIC / Passport": "IDNumber",
    "University": "University", "Nationality": "Nationality", "Residence": "Residence",
    "Unit no. ": "Unit", "Room": "Room", "Bed": "Bed", "Tenancy dates": "Tenancy Start",
    "Tenancy End": "Tenancy End", "Agreement Status": "Status", "Monthly rent": "Monthly Rent",
    "Payment frequency": "Payment Frequency", "Paid by": "PAYOR?", "Payor name": "SPONSOR NAME",
    "Copy of Tenancy Agreement": "TACopy", "Copy of Invoice": "Initial Payment Invoice Copy",
    "Total Initial Payment": "Total Initial Payment Made", "Remarks": "Remarks", "Resident type": "University",
}


def escalate(row: dict, level: str):
    """A row's status only ever gets worse, never better."""
    if RANK[level] > RANK[row["_status"]]:
        row["_status"] = level


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    # a date the team gives in the corrections file: "2027-04-30" (20 Oct 2026)
    if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", v.strip()):
        try:
            return date.fromisoformat(v.strip())
        except ValueError:
            return None
    return None


# ---------------------------------------------------------------- conversion

# DM - Unit Set Up's "3 Bedroom Apartment" is Homes' "3-bedroom"
UNIT_TYPE = {"3 BEDROOM APARTMENT": "3-bedroom", "4 BEDROOM APARTMENT": "4-bedroom",
             "1 BEDROOM APARTMENT": "1 Bedroom Apartment with attached bath", "2 BEDROOM APARTMENT": "2 Bedroom Apartment"}


def unit_types(source: Path) -> dict[str, str]:
    """unit number -> unit type, from the master's own unit list (the truth for units)."""
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if "DM - Unit Set Up" not in wb.sheetnames:
        return {}
    out = {}
    for row in wb["DM - Unit Set Up"].iter_rows(min_row=2, values_only=True):
        if row and row[1] and row[2]:
            no = re.sub(r"-(\d)A$", r"-0\1A", str(row[1]).strip().upper())
            out[no] = UNIT_TYPE.get(str(row[2]).strip().upper(), str(row[2]).strip())
    wb.close()
    return out


def unit_residences(source: Path) -> dict[str, str]:
    """unit number -> residence, from the same unit list - for a tab with no
    Residence column (the Sponsor tab, Dani 20 Oct 2026)."""
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if "DM - Unit Set Up" not in wb.sheetnames:
        return {}
    out = {}
    for row in wb["DM - Unit Set Up"].iter_rows(min_row=2, values_only=True):
        if row and len(row) > 3 and row[1] and row[3]:
            no = re.sub(r"-(\d)A$", r"-0\1A", str(row[1]).strip().upper())
            out[no] = str(row[3]).strip()
    wb.close()
    return out


def fill_next_rent(rows: list[dict], today: date):
    """Next billing period / due / amount, by the system's own rule (next-rent.ts)."""
    bun = shutil.which("bun")
    if not bun:
        return "bun not found - next rent columns left blank"
    payload = {"today": today.isoformat(), "rows": [{
        "key": str(r["Source Row"]),
        "start": r["Tenancy Start"].isoformat() if r.get("Tenancy Start") else "",
        "end": r["Tenancy End"].isoformat() if r.get("Tenancy End") else "",
        "frequency": r.get("Payment frequency") or "",
        "rent": float(r.get("Monthly rent") or 0),
        "prorated": float(r.get("First month rent - pro-rated x/x days") or 0),
        # a sponsor's advance counts in TA RENT months (see "_advance_months")
        "advance": float(r["_advance_months"] * (r.get("Monthly rent") or 0)) if r.get("_advance_months")
        else float(r.get("Advance rental (RM)") or 0),
    } for r in rows]}
    done = subprocess.run([bun, str(Path(__file__).parent / "next-rent.ts")], input=json.dumps(payload),
                          capture_output=True, text=True, cwd=ROOT)
    if done.returncode:
        return "next rent could not be worked out: " + done.stderr.strip()[:200]
    answers = json.loads(done.stdout)
    for r in rows:
        a = answers.get(str(r["Source Row"]), {})
        r["Next billing period"] = a.get("period") or ""
        r["Next payment due"] = date.fromisoformat(a["due"]) if a.get("due") else None
        r["Next payment amount"] = a.get("amount")
    return ""


# the charge columns the app adds up as the initial invoice's lines (import-plan initialLines)
INITIAL_LINE_COLUMNS = (
    "First month rent - pro-rated x/x days", "Advance rental (RM)",
    "Security Deposit (RM) (refundable) ", "Utilities deposit (RM) (refundable) ",
    "Access card deposit (refundable)", "Resident card charges", "Add on: Starter pack",
    "Admin + agreement charges",
)


def initial_balance_due(out_row: dict, paid, answer=None):
    """Owed on the initial invoice: lines - paid above RM0.01, else blank. The team's answer wins."""
    if answer is not None:
        # the team's 0 is written as 0, not blank: "checked, nothing owed" (19 Oct 2026)
        return round(float(answer), 2) if float(answer) > 0.01 else 0.0
    if paid is None:
        return None
    owed = round(sum(float(out_row.get(c) or 0) for c in INITIAL_LINE_COLUMNS) - paid, 2)
    return owed if owed > 0.01 else None


def join_notes(*parts: str) -> str:
    return " · ".join(p.strip() for p in parts if p and p.strip())


def load_corrections(path: Path | None) -> dict:
    if not path or not path.exists():
        return {"set": {}, "leaveOut": {}, "remarksHandled": {}, "applicationRows": {}, "application": {}, "nextRent": {}, "internalNotes": {}, "balanceDue": {}, "redCleared": {}, "checkedFirst": []}
    c = json.loads(path.read_text())
    fixes = {k: c.get(k, {}) for k in ("set", "leaveOut", "remarksHandled", "applicationRows", "application", "nextRent", "internalNotes", "balanceDue", "redCleared")}
    fixes["checkedFirst"] = c.get("checkedFirst", [])
    return fixes


def _nationality_words() -> dict[str, str]:
    """'MOROCCAN' or 'MOROCCO' -> 'MAR', read from the system's own lists in
    reference-data.ts (DEMONYMS and the country list) - one list, not a copy."""
    ts = (ROOT / "src/lib/reference-data.ts").read_text()
    out = {}
    for code, word in re.findall(r'^\s+([A-Z]{3}): "([^"]+)",', ts, flags=re.M):
        out[word.upper()] = code
    for code, name in re.findall(r'\{ code: "([A-Z]{3})", name: "([^"]+)"', ts):
        out[name.upper()] = code
    return out


NATIONALITY_WORDS = _nationality_words()
# 'ZWE' -> '263', from the same country list in reference-data.ts
NATIONAL_DIAL = {code: dial.lstrip("+") for code, dial in re.findall(
    r'\{ code: "([A-Z]{3})", name: "[^"]+", dial: "(\+\d+)"', (ROOT / "src/lib/reference-data.ts").read_text())}


# Same converter for every batch (Dani, 14 Oct 2026): a later tab may spell a
# column differently, or not have one Batch 1 had.
# the Sponsor tab says "Payment Status" / "Balance Due" for the initial payment
# A whole unit let to a group in one row (Dani, 22 Oct 2026: C-19-02, 4 tenants):
# "00357 00358 00359 00360" with one name, NRIC, phone and email per line. The
# TA copy / invoice copy name the lead ("TA2026_C-19-02_00358") - only they go in,
# paying the row's rent for the whole unit; the others are kept in Internal Notes.
GROUP_COPY_COLUMNS = ("TACopy", "Initial Payment Invoice Copy")
GROUP_PERSON_COLUMNS = ("StudentName", "IDNumber", "MobileNumber", "Email")


def unit_group(ids_text, *copies):
    ids = [i.zfill(5) if i.isdigit() else i for i in re.split(r"[\s,/]+", text(ids_text)) if i]
    if len(ids) < 2:
        return None
    named = [i for i in ids if any(i in text(c) for c in copies)]
    if len(named) != 1:
        return None
    return {"id": named[0], "index": ids.index(named[0]), "ids": ids}


def group_line(value, index):
    lines = [l.strip() for l in text(value).split("\n") if l.strip()]
    return lines[index] if index < len(lines) else (lines[0] if len(lines) == 1 else "")


def group_note(read, lead):
    others = []
    for n, i in enumerate(lead["ids"]):
        if i == lead["id"]:
            continue
        parts = [group_line(read(c), n) for c in GROUP_PERSON_COLUMNS]
        others.append(f"{i} " + ", ".join(p for p in parts if p))
    return "Shares the whole unit with: " + "; ".join(others)


COLUMN_ALIASES = {"Student Id": "StudentID", "Remark": "Remarks", "Payment Status": "IP Status", "Balance Due": "Balance Due on IP",
                  # the Sponsor tab (master 7): its total is what was invoiced
                  "Total Initial Payment": "Total Initial Payment Made"}
OPTIONAL_SOURCE = {"SPONSOR NAME", "IP Status", "Residence", "Tag"}
# Batch 2 onward says per row whether it goes in; only this value is converted
MIGRATE_OK = "Clean-Migrate"
CHECK_MARKS = {"CLEAN", "CHECKED", "OK"}
SHEET = "Migration Batch 1"
# Petronas pays in fixed 6-month blocks, Sep-Feb and Mar-Aug (Dani with Brachtia,
# 20 Oct 2026). For this migration: the block running now is outstanding and only
# recorded (B1-B3 Sep 2026-Feb 2027 as its own bill; B4's is their initial
# invoice), and the system invoices from the next block on.
PETRONAS_NEXT = (date(2027, 3, 1), date(2027, 8, 31))
PETRONAS_CURRENT = (date(2026, 9, 1), date(2027, 2, 28))


def day_first(d: date) -> str:
    return d.strftime("%d %b %Y")


def whole_months(a: date, b: date) -> int:
    return (b.year - a.year) * 12 + b.month - a.month + 1


def convert(source: Path, labels: Path, out: Path, report: Path, corrections: Path | None = None,
            applications: Path | None = None, batch: int = 1, sheet: str | None = None) -> dict:
    global SHEET
    # any migration tab by name (Rina, 16 Oct 2026: "Migration Batch (Non-Sponsor)")
    SHEET = sheet or f"Migration Batch {batch}"
    for p in (out, report):
        inputs = [source.resolve(), labels.resolve()] + ([applications.resolve()] if applications else [])
        if p.resolve() in inputs:
            raise SystemExit("Refusing to write over an input workbook")
    before = tuple(fingerprint(p) for p in (source, labels, applications) if p)

    headers = upload_headers(sheet3_headers(labels))
    allowed = allowed_values(labels)
    fixes = load_corrections(corrections)
    types = unit_types(source)
    residences = unit_residences(source)
    apps = application_rows(applications) if applications else []
    app_index = application_indexes(apps) if applications else None
    apps_by_row = {row["_row"]: row for row in apps}

    # Two views of the same sheet: one shows formulas and hyperlinks, the other the
    # values Excel last saved. Formulas are never evaluated here.
    wb_raw = openpyxl.load_workbook(source)
    wb_val = openpyxl.load_workbook(source, data_only=True)
    if SHEET not in wb_raw.sheetnames:
        raise SystemExit(f"The master has no '{SHEET}' tab")
    raw_ws, val_ws = wb_raw[SHEET], wb_val[SHEET]
    cols = source_columns(raw_ws)
    for alias, name in COLUMN_ALIASES.items():
        if (name, 1) not in cols and (alias, 1) in cols:
            cols[(name, 1)] = cols[(alias, 1)]

    wanted = set()
    for c in CONTRACT["columns"] + CONTRACT["migrationOnly"]:
        srcs = c.get("source") or []
        for s in [srcs] if isinstance(srcs, str) else srcs:
            wanted.add((s, c.get("sourceOccurrence", 1)))
    absent = sorted(s for s, _ in wanted if (s, 1) not in cols and s not in OPTIONAL_SOURCE)
    if absent:
        raise SystemExit(f"{SHEET} is missing expected columns: {absent}")

    cutoff = date.fromisoformat(CONTRACT["tenancyEndAfter"])
    base = source.parent

    rows, findings, formula_cells = [], [], []

    for r in range(2, raw_ws.max_row + 1):
        row_qb = text(val_ws.cell(r, cols[("StudentID", 1)]).value)
        row_qb = row_qb.zfill(5) if row_qb.isdigit() else row_qb
        # a whole unit let to a group, one row: the TA and invoice name the lead
        lead = unit_group(row_qb, *(val_ws.cell(r, cols[(c, 1)]).value for c in GROUP_COPY_COLUMNS if (c, 1) in cols))
        if lead:
            row_qb = lead["id"]
        corrected = fixes["set"].get(row_qb, {})

        def cell(name, occ=1):
            if occ == 1 and name in corrected:
                return corrected[name]
            if (name, occ) not in cols:
                return None
            value = val_ws.cell(r, cols[(name, occ)]).value
            if lead and name == "StudentID":
                return lead["id"]
            if lead and name in GROUP_PERSON_COLUMNS:
                return group_line(value, lead["index"])
            return value

        def raw_cell(name, occ=1):
            return raw_ws.cell(r, cols[(name, occ)])

        if all(val_ws.cell(r, c).value in (None, "") for c in range(1, raw_ws.max_column + 1)):
            continue

        out_row: dict[str, object] = {"Source Row": r}
        issues: list[tuple[str, str, str]] = []
        migrate = text(cell("Migrate?"))
        if migrate and migrate != MIGRATE_OK:
            # the master's own call for this row - listed in the report, not uploaded
            out_row["_left_out"] = f"Migrate? is '{migrate}' - only {MIGRATE_OK} rows go in"

        def flag(level, field, reason):
            issues.append((level, field, reason))

        qb = text(cell("StudentID"))
        # a StudentID typed as a number loses its zeros (1007 is 01007) - the
        # system pads it the same way (legacyId in src/lib/migration/sheet3.ts)
        if qb.isdigit() and len(qb) < 5:
            qb = qb.zfill(5)
            flag(WARNING, "QuickBooks ID", "StudentID was saved as a number - padded back to 5 digits")
        out_row["QuickBooks ID"] = qb
        if not qb:
            flag(BLOCKED, "QuickBooks ID", "StudentID is blank - it is the matching key")

        name = text(cell("StudentName"))
        out_row["Resident Name"] = name
        if not name:
            flag(BLOCKED, "Resident Name", "Name is blank")

        g = text(cell("Gender")).upper()
        out_row["Gender"] = GENDER.get(g, "")
        if g and g not in GENDER:
            flag(WARNING, "Gender", "Gender not recognised - left blank")

        phone_nationality = NATIONALITY_WORDS.get(text(cell("Nationality")).upper(), text(cell("Nationality")).upper())
        code, mobile, f = split_phone(text(cell("MobileNumber")), phone_nationality)
        out_row["Country Code"], out_row["Mobile Number"] = code, mobile
        if f:
            flag(f[0], "Mobile Number", f[1])

        email = text(cell("Email")).lower()
        out_row["Email"] = email
        if not email:
            flag(BLOCKED, "Email", "Email is blank")
        elif not EMAIL.match(email):
            flag(BLOCKED, "Email", "Email is not a valid address")

        nat = NATIONALITY_WORDS.get(text(cell("Nationality")).upper(), text(cell("Nationality")).upper())
        out_row["Nationality"] = nat
        nric, passport, f = split_id(nat, cell("IDNumber"))
        out_row["NRIC number"], out_row["Passport number"] = nric, passport
        if f:
            flag(f[0], "NRIC / Passport", f[1])

        uni = text(cell("University"))
        out_row["University"] = UNIVERSITY.get(uni.upper(), uni)
        # Batch 1 is students: a university means a student
        out_row["Resident type"] = "student" if uni else ""
        if app_index is not None:
            manual_row = fixes["applicationRows"].get(qb)
            if manual_row not in (None, ""):
                try:
                    manual_row = int(manual_row)
                except (TypeError, ValueError):
                    manual_row = -1
                matched = apps_by_row.get(manual_row)
                method = f"manual application row {manual_row}"
                application_finding = (
                    (WARNING, f"Application row {manual_row} was selected in batch1-corrections.json")
                    if matched else
                    (BLOCKED, f"Application row {manual_row} from batch1-corrections.json does not exist")
                )
            elif qb in fixes["application"]:
                # never sent the Google Form: the team asked them (corrections file)
                matched, method, application_finding = None, "team's answer (corrections file)", None
                enrich_from_answers(out_row, fixes["application"][qb], flag)
            else:
                matched, method, application_finding = match_application(out_row, app_index)
            out_row["Application Match"] = method
            if matched:
                enrich_from_application(out_row, matched, applications.parent, flag)
            if application_finding:
                flag(application_finding[0], "Application Form", application_finding[1])
        res = text(cell("Residence"))
        if not res and ("Residence", 1) not in cols:
            unit_no = re.sub(r"-(\d)A$", r"-0\1A", text(cell("Unit")).upper())
            res = residences.get(unit_no, "")
        out_row["Residence"] = RESIDENCE.get(res.upper(), res)
        payer = text(cell("PAYOR?")).upper()
        out_row["Paid by"] = PAID_BY.get(payer, "")
        if payer and payer not in PAID_BY:
            flag(REVIEW, "Paid by", "PAYOR? is not Self, Parent, Guardian or Sponsor")
        out_row["Payor name"] = text(cell("SPONSOR NAME")) if out_row["Paid by"] not in ("", "Self") else ""
        # a sponsor's intake batch is written in Remarks as "B3" / "B4" (Petronas):
        # it goes with the sponsor as "PETRONAS (B3)", the form sponsorLabel reads
        # (Dani, 12 Oct 2026) - and then the remark has been dealt with
        # the Sponsor tab has its own "Tag" column for the batch (Dani, 19 Oct 2026)
        tag = re.fullmatch(r"B\d{1,2}", text(cell("Tag")).upper())
        batch = tag or re.fullmatch(r"B\d{1,2}", text(cell("Remarks")).upper())
        batch_used = bool(batch and out_row["Payor name"]) and not tag
        if batch and out_row["Payor name"]:
            out_row["Payor name"] = f"{out_row['Payor name']} ({batch.group(0)})"
        if out_row["Paid by"] == "Sponsor" and not out_row["Payor name"]:
            flag(WARNING, "Payor name", "Sponsored, but SPONSOR NAME is blank")
        out_row["Unit ID"] = text(cell("UnitID"))
        # the system writes "03A" where the master list writes "3A"
        out_row["Unit no. "] = re.sub(r"-(\d)A$", r"-0\1A", text(cell("Unit")).upper())
        # the Sponsor tab writes "Room A" - the system's room is "A" (22 Oct 2026)
        out_row["Room"] = re.sub(r"^room\s+", "", text(cell("Room")), flags=re.I)
        out_row["Bed"] = text(cell("Bed"))
        unit_type = types.get(out_row["Unit no. "], "")
        if out_row["Room"].lower() == "unit" and out_row["Bed"].lower() in ("unit", ""):
            if "1 BEDROOM" in unit_type.upper():
                # a Solstice 1-bedroom has one room (A) and one bed (Single) in the
                # system - "Unit" in the master is that room (Dani, 14 Oct 2026)
                out_row["Room"], out_row["Bed"] = "A", "Single"
            else:
                # a whole-unit let: the room "Unit" with its one bed "Unit", as Homes and the
                # bulk upload make it (residents.functions.ts) - the unit must be let whole in Homes
                out_row["Room"], out_row["Bed"] = "Unit", "Unit"
                flag(WARNING, "Room", f"Whole-unit let - goes on {out_row['Unit no. '] or 'the unit'}'s Whole unit line in Homes "
                     f"(there when Website has a whole-unit rate for '{unit_type or 'its unit type'}'); the rent is this row's")
        # occupancy follows the bed: one person alone has the room as a single
        out_row["Occupancy"] = {"single": "Single", "twin 1": "Twin Share", "twin 2": "Twin Share"}.get(
            out_row["Bed"].strip().lower(), "")
        out_row["Unit Type"] = types.get(out_row["Unit no. "], "")
        if types and out_row["Unit no. "] and not out_row["Unit Type"]:
            flag(REVIEW, "Unit Type", f"{out_row['Unit no. ']} is not in DM - Unit Set Up")
        for h in ("Residence", "Unit no. ", "Room", "Bed"):
            if not out_row[h]:
                flag(BLOCKED, h, f"{h.strip()} is blank - the bed cannot be matched")

        start, end = as_date(cell("Tenancy Start")), as_date(cell("Tenancy End"))
        out_row["Tenancy Start"], out_row["Tenancy End"] = start, end
        if not start or not end:
            flag(BLOCKED, "Tenancy dates", "Tenancy start or end is not a real date")
        elif start >= end:
            flag(BLOCKED, "Tenancy dates", "Tenancy start is not before tenancy end")
        else:
            out_row["Stay Duration"] = stay_duration(start, end)
            out_row["Term"] = "Long Term" if end >= add_months(start, 6) else "Short Term"
            if end <= cutoff:
                # finished before the cutoff: not uploaded now (Darren goes in
                # Migration Batch 3 - Dani, 12 Oct 2026), only listed in the report
                out_row["_left_out"] = f"Tenancy ended {end:%d %b %Y}, on or before {cutoff:%d %b %Y} - not uploaded in this batch"

        # the existing TA's own number (Rina, 13 Oct 2026): TA<start year>_<unit>_<QB ID>
        out_row["Agreement No."] = legacy_agreement_no(start, out_row["Unit no. "], qb)

        status = text(cell("Status", 1))
        out_row["Agreement Status"] = status
        if status.lower() != "active":
            flag(REVIEW, "Agreement Status", "Status is not Active")

        rent = cell("Monthly Rent")
        out_row["Monthly rent"] = rent if isinstance(rent, (int, float)) else None
        if not isinstance(rent, (int, float)) or rent <= 0:
            flag(BLOCKED, "Monthly rent", "Monthly rent is missing or not a number")

        freq = text(cell("Payment Frequency")).upper()
        out_row["Payment frequency"] = FREQUENCY.get(freq, "")
        if freq not in FREQUENCY:
            flag(BLOCKED, "Payment frequency", "Payment frequency not recognised")

        # Unpaid (Petronas B4, Dani 19 Oct 2026): the initial invoice is out but
        # nothing is paid - there may be no copy yet, and its total is all owed
        # - unless the team has answered what is owed (corrections balanceDue wins)
        unpaid = text(cell("IP Status")).lower().startswith("unpaid") and row_qb not in fixes["balanceDue"]
        # ---- documents
        for src, header, label, required in (
            ("TACopy", "Copy of Tenancy Agreement", "Tenancy Agreement copy", True),
            ("Initial Payment Invoice Copy", "Copy of Invoice", "Invoice copy", not unpaid),
        ):
            c = raw_cell(src)
            value, f = check_document(text(cell(src)), c.hyperlink.target if c.hyperlink else None, base, label, required)
            out_row[header] = value
            if f:
                flag(f[0], header, f[1])
        # the QuickBooks invoice number is the copy's file name: "INV:IP_00894.pdf"
        # is INV:IP_00894 - kept as written, so it never clashes with the
        # system's own INV/IP/00195 numbers (Dani, 12 Oct 2026)
        inv_no = re.search(r"INV\s*[:\-]\s*IP\s*_\s*(\d+)", text(cell("Initial Payment Invoice Copy")), re.I)
        out_row["Invoice No."] = f"INV:IP_{inv_no.group(1)}" if inv_no else ""
        if not inv_no:
            flag(WARNING, "Invoice No.", "The invoice copy's name has no INV:IP number - the system will number it")
        # ---- money: read the saved value; note where it came from a formula
        def amount(src, kind):
            c = raw_cell(src)
            v = cell(src)
            if isinstance(c.value, str) and c.value.startswith("="):
                formula_cells.append((r, qb, src))
                if v is None:
                    flag(REVIEW, src, "Formula with no saved value - open and save the file in Excel")
                    return None
            if v is None or v == "":
                # Rina confirmed Migration 1 is clean: a blank charge is an
                # intentional RM0. The amount collected is still required.
                return None if kind == "paid" else 0.0
            if not isinstance(v, (int, float)):
                flag(BLOCKED, src, "Not a number")
                return None
            return float(v)

        charges = {
            "Differential Rental (QB)": amount("Differential Rental (QB)", "optional"),
            "Advance Rental": amount("Advance Rental", "core"),
            "Security Deposit": amount("Security Deposit", "core"),
            "Utility Deposit": amount("Utility Deposit", "core"),
            "Access Card Deposit": amount("Access Card Deposit", "core"),
            "Admin Charges": amount("Admin Charges", "core"),
            "Agreement Charges": amount("Agreement Charges", "optional"),
            "Personalised Tag": amount("Personalised Tag", "optional"),
        }
        paid = amount("Total Initial Payment Made", "paid")
        if unpaid:
            paid = 0.0
            flag(WARNING, "Initial Balance Due", "The master says Unpaid - imported as nothing paid, the whole initial invoice owed")
        if paid is None and not any(i[1] == "Total Initial Payment Made" for i in issues):
            flag(BLOCKED, "Total Initial Payment", "Total Initial Payment Made is blank")

        out_row["First month rent - pro-rated x/x days"] = charges["Differential Rental (QB)"]
        out_row["Advance rental (RM)"] = charges["Advance Rental"]
        # what the whole-month rule cannot read on its own (Dani, 15 Oct 2026)
        pro, adv = charges["Differential Rental (QB)"] or 0, charges["Advance Rental"] or 0
        rent_now = out_row.get("Monthly rent") or 0
        # these two signs mostly mean a renewal: the pro-rated/advance amounts are from the
        # first setup invoice, so the next payment cannot be worked out from them (16 Oct 2026)
        payor = str(out_row.get("Payor name") or "")
        # Petronas always pays half-yearly: a master that says otherwise (00968 "BI-MONTHLY")
        # made 2-month bills under a 6-month amount (Dani, 20 Oct 2026)
        if payor.upper().startswith("PETRONAS") and out_row.get("Payment frequency") not in ("semiannual", ""):
            flag(WARNING, "Payment frequency", f"Petronas pays half-yearly - the master's {text(cell('Payment Frequency'))} is taken as HALF-YEARLY")
            out_row["Payment frequency"] = "semiannual"
        if payor.upper().startswith("PETRONAS") and rent_now and isinstance(end, date) and row_qb not in fixes["nextRent"]:
            nxt_end = min(PETRONAS_NEXT[1], end)
            if nxt_end >= PETRONAS_NEXT[0]:
                fixes["nextRent"][row_qb] = {
                    "period": f"{day_first(PETRONAS_NEXT[0])} – {day_first(nxt_end)}",
                    "due": "2027-03-05",
                    "amount": round(rent_now * whole_months(PETRONAS_NEXT[0], nxt_end), 2),
                }
                out_row["_petronas"] = True
            # B1-B3: Sep 2026-Feb 2027 is outstanding, recorded without an invoice of ours
            if not payor.upper().endswith("(B4)") and end >= PETRONAS_CURRENT[0]:
                cur_end = min(PETRONAS_CURRENT[1], end)
                out_row["Current period outstanding"] = f"{day_first(PETRONAS_CURRENT[0])} – {day_first(cur_end)}"
                out_row["Current period amount"] = round(rent_now * whole_months(PETRONAS_CURRENT[0], cur_end), 2)
        # paid in full (FULLY): there is no next payment for these signs to confuse (20 Oct 2026)
        renewal = REVIEW if row_qb not in fixes["nextRent"] and out_row.get("Payment frequency") != "full" else WARNING
        if pro and isinstance(start, date) and start.day == 1:
            flag(renewal, "Differential Rental (QB)",
                 f"Starts on the 1st but has a pro-rated amount ({pro:,.2f}) - likely a renewal; ask the team for the next payment")
        # Sponsor tab: TA RENT is the rent in the resident's signed TA, since raised
        # (TA CLEAN = Problem, the TA is regenerated - Dani 19 Oct 2026). The advance
        # was billed at that old rent, so the paid months are counted by TA RENT.
        ta_rent = cell("TA RENT")
        ta_rent = float(ta_rent) if isinstance(ta_rent, (int, float)) and ta_rent > 0 else 0
        whole = lambda x, by: by and abs(x / by - round(x / by)) <= 0.01
        if adv and rent_now and ta_rent and ta_rent != rent_now and not whole(adv, rent_now) and whole(adv, ta_rent):
            out_row["_advance_months"] = round(adv / ta_rent)
            flag(WARNING, "Advance Rental",
                 f"Advance rental ({adv:,.2f}) is {round(adv / ta_rent)} months of the old TA rent ({ta_rent:,.2f}) - "
                 f"counted as {round(adv / ta_rent)} paid months; the next bills are at the current rent ({rent_now:,.2f})")
        elif adv and rent_now and abs(adv / rent_now - round(adv / rent_now)) > 0.01:
            flag(renewal, "Advance Rental",
                 f"Advance rental ({adv:,.2f}) is not whole months of rent ({rent_now:,.2f}) - likely a renewal; ask the team for the next payment")
        out_row["Security Deposit (RM) (refundable) "] = charges["Security Deposit"]
        out_row["Utilities deposit (RM) (refundable) "] = charges["Utility Deposit"]
        out_row["Access card deposit (refundable)"] = charges["Access Card Deposit"]
        out_row["Resident card charges"] = charges["Personalised Tag"]
        out_row["Add on: Starter pack"] = amount("Starter Pack", "optional")
        admin, agreement = charges["Admin Charges"], charges["Agreement Charges"]
        out_row["Admin + agreement charges"] = None if admin is None else admin + (agreement or 0)
        out_row["Total Initial Payment"] = paid
        # A QuickBooks invoice can contain extra lines (for example utilities),
        # so listed migration charges are deliberately not reconciled to its total.
        out_row["Payment Status"] = "Paid" if paid is not None else ""

        # what is still owed on the initial invoice (Dani, 19 Oct 2026): the listed
        # lines less what was paid, only above 1 sen. Paid more is not a debt and
        # not a credit. The team's answer wins - 0 clears a running-account invoice.
        # the master's own "IP Status" (Dani, 20 Oct 2026): Fully paid = checked, nothing owed
        answer = fixes["balanceDue"].get(row_qb)
        ip_status = text(cell("IP Status")).lower()
        if answer is None and ip_status.startswith("fully"):
            answer = 0
        # the master's own Balance Due on the initial payment, when it gives a number
        master_due = cell("Balance Due on IP")
        if answer is None and isinstance(master_due, (int, float)):
            answer = master_due
        out_row["Initial Balance Due"] = initial_balance_due(out_row, paid, answer)
        if ip_status.startswith("partial") and not out_row["Initial Balance Due"]:
            flag(REVIEW, "Initial Balance Due", "The master says Partially paid, but the listed lines and the amount paid show nothing owed - fill in what is owed")
        if out_row["Initial Balance Due"]:
            out_row["Payment Status"] = "Unpaid" if unpaid else "Part paid"
        # Sponsor tab: the team found the signed TA wrong - a new one is generated
        # in the system and sent for signing after import (Dani, 20 Oct 2026)
        if text(cell("TA CLEAN")).lower() == "problem":
            flag(WARNING, "TACopy", "TA CLEAN is Problem - regenerate the TA in the system and send the signing link after import")

        # Internal Notes: the team's note and the master's remarks (a sponsor's
        # batch code is not a remark) - admin only on the resident profile
        # "CLEAN" in Remarks is the team's checking mark, not a note (Dani, 19 Oct 2026)
        remark = "" if batch_used else text(cell("Remarks"))
        # "B3, Price from..." - the batch is the sponsor's, the rest is the remark
        remark = re.sub(r"^\s*B\d{1,2}\s*[,.;:-]?\s*", "", remark) if batch else remark
        if remark.strip().upper() in CHECK_MARKS:
            remark = ""
        if lead:
            remark = join_notes(remark, group_note(lambda c: val_ws.cell(r, cols[(c, 1)]).value if (c, 1) in cols else None, lead))
            flag(WARNING, "Resident Name", f"Whole unit let to {len(lead['ids'])} people in one row - {lead['id']} (named by the TA/invoice) "
                 "goes in with the unit's rent; the others are listed in Internal Notes")
        out_row["Internal Notes"] = join_notes(fixes["internalNotes"].get(row_qb, ""), remark)
        # a duplicate row is left out by its row number ("row 91"), not its ID
        left_out = fixes["leaveOut"].get(row_qb) or fixes["leaveOut"].get(f"row {r}")
        # a name the team filled red has a problem they are still checking (Dani, 19 Oct 2026)
        name_fill = raw_ws.cell(r, cols[("StudentName", 1)]).fill
        if not left_out and name_fill.fill_type and str(name_fill.fgColor.rgb).upper() == "FFFF0000":
            # the team cleared it, the red not yet taken off (corrections "redCleared", Dani 23 Oct 2026)
            if row_qb in fixes["redCleared"]:
                flag(WARNING, "Resident Name", f"Red in the master, cleared: {fixes['redCleared'][row_qb]}")
            else:
                left_out = "Name is red in the master - the team is still checking it"
        if left_out:
            out_row["_left_out"] = left_out
        if corrected:
            flag(WARNING, "Corrected", "Corrected from the team's answer: " + ", ".join(corrected) + " (batch1-corrections.json)")
        if remark and row_qb not in fixes["remarksHandled"]:
            flag(WARNING, "Remarks", "Source row has remarks (carried into Internal Notes) - read them before importing")

        # every dropdown column holds only a value on its list - else blank + flag
        for header, field in DROPDOWNS.items():
            v = out_row.get(header)
            if v in (None, "") or field not in allowed:
                continue
            if header in ("Room", "Bed") and v == "Unit":
                continue  # the whole-unit slot's own name, not a dropdown choice
            if str(v) not in allowed[field]:
                if header == "University":
                    # the system's "Other - I will type it": a typed university is kept
                    flag(WARNING, header, f"'{v}' is not on the list - kept as typed (University: Other)")
                    continue
                flag(REVIEW, header, f"'{v}' is not on the {field} list - left blank")
                out_row[header] = ""

        levels = {i[0] for i in issues}
        out_row["_status"] = BLOCKED if BLOCKED in levels else REVIEW if REVIEW in levels else WARNING if WARNING in levels else "Ready"
        rows.append(out_row)
        findings += [(r, qb, *i) for i in issues]

    # ---- left out of this batch: listed in the report, never in the upload
    left_out = [r for r in rows if r.get("_left_out")]
    rows = [r for r in rows if not r.get("_left_out")]
    findings = [f for f in findings if f[0] not in {r["Source Row"] for r in left_out}]

    # ---- duplicates across the file
    duplicates = []
    for field, level in (("QuickBooks ID", BLOCKED), ("Email", REVIEW), ("NRIC number", REVIEW), ("Passport number", REVIEW)):
        seen: dict[str, list[dict]] = {}
        for row in rows:
            v = str(row.get(field) or "").lower()
            if v:
                seen.setdefault(v, []).append(row)
        for group in seen.values():
            if len(group) > 1:
                srcs = ", ".join(str(g["Source Row"]) for g in group)
                for g in group:
                    duplicates.append((g["Source Row"], g["QuickBooks ID"], field, f"Same {field} as rows {srcs}"))
                    findings.append((g["Source Row"], g["QuickBooks ID"], level, field, f"Duplicate {field} (rows {srcs})"))
                    escalate(g, level)

    beds: dict[tuple, list[dict]] = {}
    for row in rows:
        key = (row["Residence"], row["Unit no. "], row["Room"], row["Bed"])
        if all(key) and row["Tenancy Start"] and row["Tenancy End"]:
            beds.setdefault(key, []).append(row)
    for group in beds.values():
        for i, a in enumerate(group):
            for b in group[i + 1:]:
                if a["Tenancy Start"] <= b["Tenancy End"] and b["Tenancy Start"] <= a["Tenancy End"]:
                    for x, y in ((a, b), (b, a)):
                        duplicates.append((x["Source Row"], x["QuickBooks ID"], "Bed", f"Same bed, overlapping dates with row {y['Source Row']}"))
                        findings.append((x["Source Row"], x["QuickBooks ID"], BLOCKED, "Bed", f"Bed double-booked with row {y['Source Row']}"))
                        escalate(x, BLOCKED)

    note = fill_next_rent(rows, date.today())
    if note:
        print(f"Note: {note}")
    # Brachtia admin checked the next payment: their figures win (Dani, 15 Oct 2026)
    for r in rows:
        admin = fixes["nextRent"].get(r.get("QuickBooks ID"))
        if not admin:
            continue
        changed = set()
        if admin.get("due"):
            due = date.fromisoformat(admin["due"])
            if due != r.get("Next payment due"):
                changed.add("Next payment due")
            r["Next payment due"] = due
        if admin.get("amount") not in (None, ""):
            amount = float(admin["amount"])
            if r.get("Next payment amount") is None or abs(amount - float(r["Next payment amount"])) > 0.005:
                changed.add("Next payment amount")
            r["Next payment amount"] = amount
        # the cells admin changed are marked in the upload file, to see the difference
        r["_admin_changed"] = changed
        if admin.get("period"):
            if admin["period"] != r.get("Next billing period"):
                changed.add("Next billing period")
            r["Next billing period"] = admin["period"]
    group = (re.search(r"\((.+)\)", sheet).group(1) if "(" in sheet else sheet) if sheet else ""
    # residents the team already checked come first, in their order, then a line (Dani, 16 Oct 2026)
    first = {qb: i for i, qb in enumerate(fixes["checkedFirst"])}
    rows.sort(key=lambda r: first.get(r.get("QuickBooks ID"), len(first)))
    checked = sum(1 for r in rows if r.get("QuickBooks ID") in first and not r.get("_left_out"))
    write_upload(out, headers, rows, group, checked)
    letters = {name: get_column_letter(col) for (name, occ), col in cols.items() if occ == 1}
    where = lambda field: (f"{letters[FIELD_SOURCE.get(field, field)]} · {FIELD_SOURCE.get(field, field)}"
                           if FIELD_SOURCE.get(field, field) in letters else field)
    write_report(report, headers, rows, findings, duplicates, formula_cells, where, left_out)

    after = tuple(fingerprint(p) for p in (source, labels, applications) if p)
    if after != before:
        raise SystemExit("An input workbook changed during conversion - this should never happen")

    counts = {s: sum(1 for r in rows if r["_status"] == s) for s in ("Ready", WARNING, REVIEW, BLOCKED)}
    return {"rows": len(rows), "left_out": len(left_out), **counts}


# ---------------------------------------------------------------- output

# The look of Data Labelling: Arial, the header bold on light yellow (12 Oct 2026)
YELLOW = PatternFill("solid", fgColor="FFFFF2CC")
ARIAL, ARIAL_BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)


def tidy(ws, header_row=1, wrap_cols=(), max_width=60):
    """Header bold on yellow, every cell Arial, columns as wide as what is in them."""
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=header_row):
        for c in row:
            if c.value is None:
                continue
            c.font = ARIAL_BOLD if c.row == header_row else ARIAL
            longest = max(len(line) for line in str(c.value).split("\n")) if not hasattr(c.value, "year") else 11
            widths[c.column] = max(widths.get(c.column, 0), longest)
    for c in ws[header_row]:
        c.fill = YELLOW
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(w + 3, 10), max_width)
    for col in wrap_cols:
        for row in ws.iter_rows(min_row=header_row + 1, min_col=col, max_col=col):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    # short codes (row numbers, IDs, statuses) sit in the middle of their column
    for c in ws[header_row]:
        if str(c.value) in ("Source Row", "QuickBooks ID", "Count", "Status", "Level", "Tenancy Agreement", "Invoice"):
            for row in ws.iter_rows(min_row=header_row, min_col=c.column, max_col=c.column):
                for cell in row:
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=cell.row == header_row)
                    if str(c.value) == "QuickBooks ID":
                        cell.number_format = "@"
    ws.row_dimensions[header_row].height = 30
    ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate


DATE_HEADERS = {"Tenancy Start", "Tenancy End", "Agreement Date", "Invoice Date", "Invoice Due Date", "Next payment due"}
TEXT_HEADERS = {"QuickBooks ID", "Mobile Number", "Country Code", "NRIC number", "Passport number", "Unit ID", "Unit no. ", "Room"}
DATE_KEYS = {key(h) for h in DATE_HEADERS}
TEXT_KEYS = {key(h) for h in TEXT_HEADERS}


# a cell Brachtia admin changed: light blue, apart from a blocked row's light red
ADMIN_FILL = PatternFill("solid", fgColor="FFDBEAFE")
ROW_FILL = {
    BLOCKED: PatternFill("solid", fgColor="FFFCE4E4"),
    REVIEW: PatternFill("solid", fgColor="FFFFE8CC"),
}


def write_upload(path: Path, headers: list[str], rows: list[dict], group: str = "", checked: int = 0):
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    # the upload reads the first tab whatever its name, so the tab can say which group it is
    ws.title = f"{CONTRACT['targetSheet']} - {group}"[:31] if group else CONTRACT["targetSheet"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in rows:
        by_key = {key(k): v for k, v in row.items()}
        ws.append([by_key.get(key(h)) for h in headers])
        for c in ws[ws.max_row]:
            h = key(headers[c.column - 1])
            if h in DATE_KEYS and c.value:
                c.number_format = "dd-mmm-yyyy"
            elif h in TEXT_KEYS:
                c.number_format = "@"  # stays text, so a leading zero survives
    wrap_headers = {
        "Address", "Medical Condition - Remarks", "Photograph", "NRIC / Passport Copy",
        "Offer Letter", "Application Match",
    }
    wrap_cols = [i for i, header in enumerate(headers, start=1) if header in wrap_headers]
    tidy(ws, wrap_cols=wrap_cols, max_width=34)
    for row_no in range(2, ws.max_row + 1):
        ws.row_dimensions[row_no].height = 45
    if checked:
        line = Side(style="thick", color="FF000000")
        for c in ws[checked + 1]:
            c.border = Border(bottom=line)
    # the row's check result at a glance (Dani, 14 Oct 2026): light red = Blocked,
    # light amber = Needs review; the upload ignores colours
    for row_no, row in enumerate(rows, start=2):
        fill = ROW_FILL.get(row.get("_status"))
        if fill:
            for c in ws[row_no]:
                c.fill = fill
        for c in ws[row_no]:
            if headers[c.column - 1] in row.get("_admin_changed", ()):
                c.fill = ADMIN_FILL
    ws.freeze_panes = "B2"
    wb.save(path)


def write_report(path: Path, headers, rows, findings, duplicates, formula_cells, where=lambda f: f, left_out=()):
    wb = openpyxl.Workbook()
    fills = {"Ready": "C6EFCE", WARNING: "FFEB9C", REVIEW: "FFD8A8", BLOCKED: "FFC7CE"}

    def sheet(title, head, data, status_col=None, note=""):
        ws = wb.create_sheet(title)
        # a title and a one-line note above the table, like Data Labelling's Sheet2
        ws.append([title])
        ws.append([note])
        ws.append(head)
        for d in data:
            ws.append(list(d))
            if status_col is not None and d[status_col] in fills:
                ws.cell(ws.max_row, status_col + 1).fill = PatternFill("solid", fgColor=fills[d[status_col]])
        last = len(head)
        tidy(ws, header_row=3, wrap_cols=[last] if last > 2 else [], max_width=70)
        ws["A1"].font = Font(name="Arial", size=12, bold=True)
        ws["A2"].font = Font(name="Arial", size=10, italic=True, color="FF666666")
        ws["A1"].fill = ws["A2"].fill = PatternFill(fill_type=None)
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 14)
        return ws

    wb.remove(wb.active)
    counts = {s: sum(1 for r in rows if r["_status"] == s) for s in ("Ready", WARNING, REVIEW, BLOCKED)}
    sheet("Validation Summary", ["Measure", "Count"], [("Rows converted", len(rows)), *counts.items(),
          ("Cells read from a saved formula result", len(formula_cells))],
          note="How many rows are ready. Green = Ready, yellow = Warning, orange = Needs review, red = Blocked.")
    summary = wb["Validation Summary"]
    for row in summary.iter_rows(min_row=4):
        if row[0].value in fills:
            row[0].fill = PatternFill("solid", fgColor=fills[row[0].value])
    def fix_in(source_row):
        seen = []
        for f in findings:
            if f[0] == source_row and where(f[3]) not in seen:
                seen.append(where(f[3]))
        return "\n".join(seen)

    sheet("Row Status", ["Source Row", "QuickBooks ID", "Name", "Status", "Fix in column", "Reasons"],
          [(r["Source Row"], r["QuickBooks ID"], r.get("Resident Name"), r["_status"], fix_in(r["Source Row"]),
            "\n".join(f[4] for f in findings if f[0] == r["Source Row"])) for r in rows], status_col=3,
          note="Every resident. Source Row and 'Fix in column' point at the cell in " + SHEET + ".")
    sheet("Left out", ["Source Row", "QuickBooks ID", "Name", "Why"],
          [(r["Source Row"], r["QuickBooks ID"], r.get("Resident Name"), r["_left_out"]) for r in left_out],
          note="In " + SHEET + " but not in the upload file. They are handled in a later batch.")
    sheet("Missing Required", ["Source Row", "QuickBooks ID", "Field", "Reason"],
          [(f[0], f[1], f[3], f[4]) for f in findings if f[2] == BLOCKED and ("blank" in f[4] or "missing" in f[4])])
    names = {r["Source Row"]: r.get("Resident Name") for r in rows}
    order = {BLOCKED: 0, REVIEW: 1, WARNING: 2}
    sheet("Manual Review", ["Source Row", "QuickBooks ID", "Name", "Level", "Fix in column", "Reason"],
          [(f[0], f[1], names.get(f[0]), f[2], where(f[3]), f[4])
           for f in sorted(findings, key=lambda f: (order.get(f[2], 9), f[0]))], status_col=3,
          note="Every problem, worst first. Fix it in that column of " + SHEET + " (references folder), save, run again.")
    sheet("Duplicates", ["Source Row", "QuickBooks ID", "Field", "Detail"], duplicates)
    application_fields = {
        "Application Form", "Referral Source", "Marital status", "Race", "Religion", "Address",
        "Postal Code", "State / Region / Province", "Country", "Medical Condition",
        "Medical Condition - Remarks", "Photograph", "NRIC / Passport Copy", "Offer Letter",
    }
    def application_status(row):
        if not row.get("Application Source Row"):
            return BLOCKED
        return WARNING if any(f[0] == row["Source Row"] and f[3] in application_fields for f in findings) else "Ready"

    sheet("Application Matches",
          ["Source Row", "QuickBooks ID", "Name", "Application Source Row", "Matched by", "Status", "Reasons"],
          [(r["Source Row"], r["QuickBooks ID"], r.get("Resident Name"), r.get("Application Source Row"),
            r.get("Application Match"), application_status(r),
            "\n".join(f[4] for f in findings if f[0] == r["Source Row"] and f[3] in application_fields))
           for r in rows], status_col=5,
          note="How Batch 1 was connected to gform responses. A match uses NRIC/passport, university Student ID, email, phone or normalized name; conflicts stay visible.")
    sheet("Documents", ["Source Row", "QuickBooks ID", "Tenancy Agreement", "Invoice", "Photograph",
                        "NRIC / Passport Copy", "Offer Letter", "Problems"],
          [(r["Source Row"], r["QuickBooks ID"],
            "OK" if r.get("Copy of Tenancy Agreement") and not any(f[0] == r["Source Row"] and f[3] == "Copy of Tenancy Agreement" for f in findings) else "Problem",
            "OK" if r.get("Copy of Invoice") and not any(f[0] == r["Source Row"] and f[3] == "Copy of Invoice" for f in findings) else "Problem",
            "OK" if r.get("Photograph") else "Problem",
            "OK" if r.get("NRIC / Passport Copy") else "Problem",
            "OK" if r.get("Offer Letter") else "Problem",
            "; ".join(f[4] for f in findings if f[0] == r["Source Row"] and
                      (f[3].startswith("Copy of") or f[3] in {"Application Form", "Photograph", "NRIC / Passport Copy", "Offer Letter"})))
           for r in rows])
    sheet("Formula Cells", ["Source Row", "QuickBooks ID", "Column"], formula_cells)
    def mapped_source(column):
        source = column.get("source") or column.get("applicationSource") or []
        value = source if isinstance(source, str) else " + ".join(source)
        return f"Application: {value}" if column.get("applicationSource") else value

    mapping = [(c["sheet3"], c["key"], mapped_source(c),
                c.get("db") or "", "Required" if c.get("required") else "Optional", c["kind"],
                c.get("transform", ""), c.get("validation", ""), c.get("notes", ""))
               for c in CONTRACT["columns"] + CONTRACT["migrationOnly"]]
    mapping += [("(not carried)", "", src, "", "", "unmapped", "", "", why) for src, why in CONTRACT["unmapped"].items()]
    sheet("Mapping", ["Sheet3 header", "System key", "Source header", "Database", "Required", "Kind", "Transform", "Validation", "Notes"], mapping)
    wb.save(path)


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[1])
    p.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    p.add_argument("--labels", type=Path, default=DEFAULT_LABELS)
    p.add_argument("--applications", type=Path, default=DEFAULT_APPLICATIONS)
    p.add_argument("--batch", type=int, default=1, help="which 'Migration Batch N' tab to convert")
    p.add_argument("--sheet", help="a migration tab by its name, e.g. 'Migration Batch (Non-Sponsor)'")
    p.add_argument("--out", type=Path)
    p.add_argument("--report", type=Path)
    p.add_argument("--corrections", type=Path)
    a = p.parse_args(argv)
    n = a.batch
    # "Migration Batch (Non-Sponsor)" -> "Non-Sponsor"; a numbered batch -> "Batch N"
    name = (re.search(r"\((.+)\)", a.sheet).group(1) if a.sheet and "(" in a.sheet else a.sheet) if a.sheet else f"Batch {n}"
    slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
    a.out = a.out or ROOT / f"migration/residents/output/Data Labelling - {name} Prepared.xlsx"
    a.report = a.report or ROOT / f"migration/residents/output/{name} Conversion Report.xlsx"
    if a.corrections is None:
        own = ROOT / f"references/{'batch' + str(n) if not a.sheet else slug}-corrections.json"
        a.corrections = own if own.exists() else None
    result = convert(a.source, a.labels, a.out, a.report, a.corrections, a.applications, n, a.sheet)
    print(f"Converted {result['rows']} rows: {result['Ready']} Ready, {result[WARNING]} Warning, "
          f"{result[REVIEW]} Needs review, {result[BLOCKED]} Blocked · {result['left_out']} left out (see the Left out tab)")
    print(f"Upload file: {a.out.relative_to(ROOT) if a.out.is_relative_to(ROOT) else a.out}")
    print(f"Report:      {a.report.relative_to(ROOT) if a.report.is_relative_to(ROOT) else a.report}")


if __name__ == "__main__":
    sys.exit(main())
