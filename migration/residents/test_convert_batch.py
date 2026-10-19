"""
Tests for the Batch 1 converter, on made-up residents only.

Run:
    python3 -m unittest migration/residents/test_convert_batch.py
"""

import json
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).parent))
import convert_batch as cb  # noqa: E402

SOURCE_HEADERS = [
    "Residence", "UnitID", "Unit", "Room", "Bed", "StudentID", "StudentName", "Remarks", "Gender",
    "University", "Nationality", "PAYOR?", "SPONSOR NAME", "IDNumber", "MobileNumber", "Email", "TACopy",
    "Tenancy Start", "Tenancy End", "Payment Frequency", "Duration", "Contracting Entity", "Status",
    "Monthly Rent", "Differential Rental (QB)", "Advance Rental", "Security Deposit", "Utility Deposit",
    "Access Card Deposit", "Admin Charges", "Agreement Charges", "Personalised Tag", "Car Park",
    "Starter Pack", "Total Initial Payment Made", "Initial Payment Invoice Copy", "Status",
]


def resident(**over):
    """A synthetic Migration 1 resident. Invoice totals may contain extra lines."""
    r = {
        "Residence": "Test Residence", "UnitID": "U900", "Unit": "Z-01-01", "Room": "A", "Bed": "Single",
        "StudentID": "90001", "StudentName": "Test Person One", "Gender": "F", "University": "Test Uni",
        "Nationality": "IND", "PAYOR?": "SELF", "IDNumber": "Z1234567", "MobileNumber": "(+91)9000000001",
        "Email": "one@example.test", "TACopy": ("TA", "https://drive.google.com/file/d/fake-ta-1"),
        "Tenancy Start": datetime(2026, 9, 14), "Tenancy End": datetime(2027, 6, 30),
        "Payment Frequency": "MONTHLY", "Duration": "=1+1", "Contracting Entity": "Test PLT", "Status": "Active",
        "Monthly Rent": 800, "Differential Rental (QB)": 100, "Advance Rental": 800, "Security Deposit": 800,
        "Utility Deposit": 200, "Access Card Deposit": 50, "Admin Charges": 100, "Agreement Charges": 50,
        "Personalised Tag": 30, "Total Initial Payment Made": 2130,
        "Initial Payment Invoice Copy": ("INV:IP_90001.pdf", "https://drive.google.com/file/d/fake-inv-1"),
        "Status#2": True,
    }
    r.update(over)
    return r


class Fixture:
    def __init__(self, rows, tmp: Path):
        self.source, self.labels = tmp / "master.xlsx", tmp / "labels.xlsx"
        self.out, self.report = tmp / "out" / "prepared.xlsx", tmp / "out" / "report.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Migration Batch 1"
        ws.append(SOURCE_HEADERS)
        for i, r in enumerate(rows, start=2):
            for c, h in enumerate(SOURCE_HEADERS, start=1):
                key = "Status#2" if c == len(SOURCE_HEADERS) else h
                v = r.get(key)
                if isinstance(v, tuple):
                    ws.cell(i, c, v[0]).hyperlink = v[1]
                else:
                    ws.cell(i, c, v)
        wb.save(self.source)

        # Sheet3 from the real labelling workbook's header list, with no data rows
        lab = openpyxl.Workbook()
        lab.active.title = "Sheet3"
        lab.active.append([c["sheet3"] for c in cb.CONTRACT["columns"] if not c.get("addedToSheet3")])
        lab.save(self.labels)

    def run(self, corrections=None):
        result = cb.convert(self.source, self.labels, self.out, self.report, corrections)
        ws = openpyxl.load_workbook(self.out)["Sheet3"]
        head = [c.value for c in ws[1]]
        rows = [dict(zip(head, (c.value for c in row))) for row in ws.iter_rows(min_row=2)]
        rep = openpyxl.load_workbook(self.report)
        status = {r[0]: r for r in rep["Row Status"].iter_rows(min_row=2, values_only=True)}
        return result, head, rows, status, ws


class ConverterTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def convert(self, *rows):
        return Fixture(list(rows), self.tmp).run()

    def test_clean_row_is_ready(self):
        result, head, rows, status, _ = self.convert(resident())
        self.assertEqual(status[2][3], "Ready")
        self.assertEqual(rows[0]["Payment Status"], "Paid")

    def test_no_columns_are_added_to_the_sheet3_template(self):
        _, head, *_ = self.convert(resident())
        self.assertNotIn("Copy of Access Card Form", head)
        self.assertIn("Bed", head)
        self.assertIn("Photograph", head)
        self.assertIn("NRIC / Passport Copy", head)
        self.assertIn("Offer Letter", head)

    def test_dates_stay_real_dates_and_ids_stay_text(self):
        _, _, rows, _, ws = self.convert(resident(StudentID="00042"))
        self.assertIsInstance(rows[0]["Tenancy Start"], datetime)
        self.assertEqual(rows[0]["QuickBooks ID"], "00042")

    def test_student_id_saved_as_a_number_gets_its_zeros_back(self):
        _, _, rows, status, _ = self.convert(resident(StudentID=1007))
        self.assertEqual(rows[0]["QuickBooks ID"], "01007")
        self.assertEqual(status[2][1], "01007")

    def test_foreigner_with_a_malaysian_601_number(self):
        _, _, rows, _, _ = self.convert(resident(Nationality="IND", MobileNumber="60123456789"))
        self.assertEqual((rows[0]["Country code"], rows[0]["Mobile number"]), ("+60", "123456789"))

    def test_phone_is_split_into_country_code_and_number(self):
        _, _, rows, _, _ = self.convert(resident(MobileNumber="(+91) 900 000 0001"))
        self.assertEqual((rows[0]["Country code"], rows[0]["Mobile number"]), ("+91", "9000000001"))

    def test_local_malaysian_number_gets_60_with_a_warning(self):
        _, _, rows, status, _ = self.convert(resident(MobileNumber="012-3456789"))
        self.assertEqual((rows[0]["Country code"], rows[0]["Mobile number"]), ("+60", "123456789"))
        self.assertIn("taken as Malaysian", status[2][5])

    def test_malaysian_number_that_lost_its_leading_zero_gets_60(self):
        _, _, rows, status, _ = self.convert(resident(MobileNumber="199813238", Nationality="MYS"))
        self.assertEqual((rows[0]["Country code"], rows[0]["Mobile number"]), ("+60", "199813238"))
        self.assertIn("lost its leading zero", status[2][5])

    def test_malaysian_number_that_starts_with_60_is_split(self):
        _, _, rows, status, _ = self.convert(resident(MobileNumber="601124339518", Nationality="MYS"))
        self.assertEqual((rows[0]["Country code"], rows[0]["Mobile number"]), ("+60", "1124339518"))
        self.assertIn("split into +60", status[2][5])

    def test_malaysian_gets_nric_and_foreigner_gets_passport(self):
        _, _, rows, _, _ = self.convert(
            resident(Nationality="MYS", IDNumber="010203-04-0506"),
            resident(StudentID="90002", Email="two@example.test", Room="B"),
        )
        self.assertEqual((rows[0]["NRIC number"], rows[0]["Passport number"]), ("010203-04-0506", None))
        self.assertEqual((rows[1]["NRIC number"], rows[1]["Passport number"]), (None, "Z1234567"))

    def test_nric_that_lost_its_zero_gets_it_back(self):
        # 010203-04-0506 saved as a number is 10203040506 (Dani, 12 Oct 2026)
        _, _, rows, status, _ = self.convert(resident(Nationality="MYS", IDNumber=10203040506, MobileNumber="+60123456789"))
        self.assertEqual(rows[0]["NRIC number"], "010203-04-0506")
        self.assertEqual(status[2][3], "Warning")

    def test_nric_whose_padded_digits_are_not_a_date_still_needs_review(self):
        _, _, rows, status, _ = self.convert(resident(Nationality="MYS", IDNumber=19903040506, MobileNumber="+60123456789"))
        self.assertIsNone(rows[0]["NRIC number"])
        self.assertEqual(status[2][3], "Needs review")

    def test_missing_nationality_needs_review(self):
        _, _, _, status, _ = self.convert(resident(Nationality=None))
        self.assertIn("ID type cannot be told", status[2][5])

    def test_invoice_total_is_not_reconciled_to_the_listed_lines(self):
        # an invoice with extra lines (utilities) paid more than the lines listed:
        # still Paid and Ready. Paid LESS is now owed (19 Oct 2026, test below).
        _, _, rows, status, _ = self.convert(resident(**{"Total Initial Payment Made": 2300}))
        self.assertEqual(rows[0]["Payment Status"], "Paid")
        self.assertEqual(status[2][3], "Ready")

    def test_paid_less_than_the_lines_owes_the_rest(self):
        # 00991-shaped, made-up numbers: lines 2,130.00, paid 1,998.29 -> owes 131.71
        _, _, rows, status, _ = self.convert(resident(**{"Total Initial Payment Made": 1998.29}))
        self.assertEqual(rows[0]["Initial Balance Due"], 131.71)
        self.assertEqual(rows[0]["Payment Status"], "Part paid")
        self.assertEqual(status[2][3], "Ready")

    def test_paid_more_than_the_lines_is_not_a_debt(self):
        # 00952 / 00993: paid a little more than the listed lines
        _, _, rows, _, _ = self.convert(resident(**{"Total Initial Payment Made": 2141.16}))
        self.assertIsNone(rows[0]["Initial Balance Due"])
        self.assertEqual(rows[0]["Payment Status"], "Paid")

    def test_team_clears_a_running_account_and_sets_a_known_debt(self):
        # 00721: extra lines on a running account - the team says nothing is owed;
        # 00994: the team knows what is owed (57.94)
        fx = Fixture([resident(**{"Total Initial Payment Made": 2000}),
                      resident(StudentID="90002", Email="two@example.test", Unit="Z-01-02")], self.tmp)
        corr = self.tmp / "corr.json"
        corr.write_text(json.dumps({"balanceDue": {"90001": 0, "90002": 57.94}}))
        _, _, rows, _, _ = fx.run(corr)
        by = {r["QuickBooks ID"]: r for r in rows}
        self.assertIsNone(by["90001"]["Initial Balance Due"])
        self.assertEqual(by["90001"]["Payment Status"], "Paid")
        self.assertEqual(by["90002"]["Initial Balance Due"], 57.94)
        self.assertEqual(by["90002"]["Payment Status"], "Part paid")

    def test_internal_notes_join_the_team_note_and_the_remarks(self):
        fx = Fixture([resident(Remarks="Late key return"),
                      resident(StudentID="90002", Email="two@example.test", Unit="Z-01-02", Remarks="Prefers WhatsApp"),
                      resident(StudentID="90003", Email="three@example.test", Unit="Z-01-03")], self.tmp)
        corr = self.tmp / "corr.json"
        corr.write_text(json.dumps({"internalNotes": {"90001": "Balance due RM57.94 (Rina, 16 Oct)", "90003": "Only the team's"}}))
        _, _, rows, _, _ = fx.run(corr)
        by = {r["QuickBooks ID"]: r["Internal Notes"] for r in rows}
        self.assertEqual(by["90001"], "Balance due RM57.94 (Rina, 16 Oct) · Late key return")
        self.assertEqual(by["90002"], "Prefers WhatsApp")
        self.assertEqual(by["90003"], "Only the team's")

    def test_whole_unit_let_keeps_the_unit_slot(self):
        _, _, rows, _, _ = self.convert(resident(Room="Unit", Bed="Unit"))
        self.assertEqual((rows[0]["Room"], rows[0]["Bed"]), ("Unit", "Unit"))

    def test_clean_checking_mark_is_not_a_note(self):
        _, _, rows, _, _ = self.convert(resident(Remarks="CLEAN"))
        self.assertIn(rows[0]["Internal Notes"], (None, ""))

    def test_no_note_and_no_remarks_is_empty(self):
        _, _, rows, _, _ = self.convert(resident())
        self.assertIn(rows[0]["Internal Notes"], (None, ""))

    def test_payment_within_a_cent_is_paid(self):
        _, _, rows, _, _ = self.convert(resident(**{"Total Initial Payment Made": 2130.004}))
        self.assertEqual(rows[0]["Payment Status"], "Paid")

    def test_blank_charge_is_intentionally_zero(self):
        _, _, rows, status, _ = self.convert(resident(**{"Advance Rental": None, "Total Initial Payment Made": 2130}))
        self.assertEqual(rows[0]["Advance rental (RM)"], 0)
        self.assertEqual(rows[0]["Payment Status"], "Paid")
        self.assertEqual(status[2][3], "Ready")

    def test_blank_optional_charge_counts_as_not_charged(self):
        _, _, rows, _, _ = self.convert(resident(**{"Personalised Tag": None, "Total Initial Payment Made": 2100}))
        self.assertEqual(rows[0]["Payment Status"], "Paid")

    def test_tenancy_ending_by_30_september_2026_is_left_out(self):
        # not uploaded now - Darren goes in Migration Batch 3 (Dani, 12 Oct 2026)
        result, _, rows, status, _ = self.convert(resident(**{"Tenancy End": datetime(2026, 9, 30)}))
        self.assertEqual((result["rows"], result["left_out"], rows), (0, 1, []))
        left = openpyxl.load_workbook(self.tmp / "out" / "report.xlsx")["Left out"]
        self.assertEqual(left.cell(4, 2).value, "90001")

    def test_tenancy_ending_in_october_stays_for_the_checkout_test(self):
        # Dani, 12 Oct 2026: October endings stay, to test their checkout statements
        _, _, _, status, _ = self.convert(resident(**{"Tenancy End": datetime(2026, 10, 31)}))
        self.assertEqual(status[2][3], "Ready")

    def test_start_after_end_is_blocked(self):
        _, _, _, status, _ = self.convert(resident(**{"Tenancy Start": datetime(2027, 7, 1)}))
        self.assertEqual(status[2][3], "Blocked")

    def test_not_active_needs_review_and_the_second_status_column_is_ignored(self):
        _, _, rows, status, _ = self.convert(resident(Status="Booked", **{"Status#2": "Active"}))
        self.assertEqual(rows[0]["Agreement Status"], "Booked")
        self.assertEqual(status[2][3], "Needs review")

    def test_duplicate_student_id_blocks_both_rows(self):
        _, _, _, status, _ = self.convert(resident(), resident(Email="two@example.test", Room="B"))
        self.assertEqual((status[2][3], status[3][3]), ("Blocked", "Blocked"))

    def test_same_bed_with_overlapping_dates_is_blocked(self):
        _, _, _, status, _ = self.convert(resident(), resident(StudentID="90002", Email="two@example.test", IDNumber="Z7654321"))
        self.assertEqual(status[3][3], "Blocked")
        self.assertIn("double-booked", status[3][5])

    def test_missing_invoice_copy_is_blocked(self):
        _, _, _, status, _ = self.convert(resident(**{"Initial Payment Invoice Copy": None}))
        self.assertEqual(status[2][3], "Blocked")

    def test_link_target_is_used_not_the_display_text(self):
        _, _, rows, _, _ = self.convert(resident())
        self.assertEqual(rows[0]["Copy of Tenancy Agreement"], "https://drive.google.com/file/d/fake-ta-1")

    def test_local_document_must_exist(self):
        (self.tmp / "ta.pdf").write_bytes(b"%PDF")
        _, _, _, status, _ = self.convert(resident(TACopy="ta.pdf", **{"Initial Payment Invoice Copy": "missing.pdf"}))
        self.assertIn("local file that does not exist", status[2][5])
        self.assertNotIn("Tenancy Agreement copy", status[2][5])

    def test_formula_is_not_run(self):
        # openpyxl writes no saved result for a formula, like a file never opened in Excel
        _, _, rows, status, _ = self.convert(resident(**{"Advance Rental": "=400*2"}))
        self.assertIsNone(rows[0]["Advance rental (RM)"])
        self.assertIn("Formula with no saved value", status[2][5])

    def test_stay_duration_matches_datedif(self):
        self.assertEqual(cb.stay_duration(date(2026, 9, 1), date(2027, 6, 30)), "10 Months 0 Days")
        self.assertEqual(cb.stay_duration(date(2026, 9, 14), date(2027, 6, 30)), "9 Months 17 Days")

    def test_source_workbook_is_unchanged(self):
        f = Fixture([resident()], self.tmp)
        before = cb.fingerprint(f.source)
        f.run()
        self.assertEqual(cb.fingerprint(f.source), before)

    def test_refuses_to_overwrite_an_input(self):
        f = Fixture([resident()], self.tmp)
        with self.assertRaises(SystemExit):
            cb.convert(f.source, f.labels, f.source, f.report)


if __name__ == "__main__":
    unittest.main()


ALLOWED = {
    "Gender": ["Male", "Female"], "Nationality": ["IND", "MYS"], "Resident type": ["student", "employed"],
    "University (application form)": ["MMU", "HWUM", "UoC", "CU", "Other"],
    "Residence": ["The Arc, Cyberjaya", "Solstice Residence, Cyberjaya"],
    "Payment frequency": ["monthly", "bimonthly", "quarterly", "semiannual", "annual", "full"],
    "Paid by": ["Parent", "Guardian", "Sponsor", "Self"], "Room": ["A", "B", "C", "D"], "Bed": ["Single", "Twin 1", "Twin 2"],
}


class ValueMapTests(unittest.TestCase):
    """Batch 1's spellings become the system's values (the Value map tab, 12 Oct 2026)."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def convert(self, *rows):
        f = Fixture(list(rows), self.tmp)
        lab = openpyxl.load_workbook(f.labels)
        ws = lab.create_sheet("Allowed values")
        ws.append(["Allowed values"]); ws.append([""]); ws.append(["Field", "Stored value", "Shown as", "Where"])
        for field, values in ALLOWED.items():
            for v in values:
                ws.append([field, v, v, "test"])
        lab.save(f.labels)
        return f.run()

    def test_values_become_the_system_values(self):
        _, _, rows, status, _ = self.convert(resident(**{
            "Residence": "The Arc, Cyberjaya", "University": "UOC", "Payment Frequency": "HALF-YEARLY", "Unit": "B-08-3A"}))
        r = rows[0]
        self.assertEqual((r["University"], r["Payment frequency"], r["Resident type"], r["Unit no. "]),
                         ("UoC", "semiannual", "student", "B-08-03A"))
        self.assertEqual(status[2][3], "Ready")

    def test_bi_monthly_and_fully(self):
        _, _, rows, _, _ = self.convert(resident(**{"Residence": "ARC", "Payment Frequency": "BI-MONTHLY"}), resident(**{
            "Residence": "ARC", "Payment Frequency": "FULLY", "StudentID": "90002", "Email": "two@example.test", "Bed": "Twin 1"}))
        self.assertEqual([r["Payment frequency"] for r in rows], ["bimonthly", "full"])
        self.assertEqual(rows[0]["Residence"], "The Arc, Cyberjaya")

    def test_payor(self):
        _, _, rows, _, _ = self.convert(
            resident(**{"Residence": "ARC", "PAYOR?": "SPONSOR", "SPONSOR NAME": "Test Sponsor Bhd"}),
            resident(**{"Residence": "ARC", "PAYOR?": "SELF", "SPONSOR NAME": "ignored", "StudentID": "90002",
                        "Email": "two@example.test", "Bed": "Twin 1"}))
        self.assertEqual((rows[0]["Paid by"], rows[0]["Payor name"]), ("Sponsor", "Test Sponsor Bhd"))
        self.assertEqual((rows[1]["Paid by"], rows[1]["Payor name"]), ("Self", None))

    def test_a_value_not_on_the_list_is_blank_and_flagged_never_added(self):
        _, _, rows, status, _ = self.convert(resident(**{"Residence": "ARC", "Gender": "X"}))
        self.assertIn(rows[0]["Gender"], (None, ""))
        self.assertNotEqual(status[2][3], "Ready")

    def test_an_unlisted_university_is_kept_as_typed(self):
        # the system's own "Other - I will type it" keeps a typed university (14 Oct 2026)
        _, _, rows, status, _ = self.convert(resident(**{"Residence": "ARC", "University": "Some New College"}))
        self.assertEqual(rows[0]["University"], "Some New College")
        self.assertEqual(status[2][3], "Warning")


class SponsorBatchTests(ValueMapTests):
    def test_b3_in_remarks_goes_with_the_sponsor(self):
        _, _, rows, status, _ = self.convert(resident(**{
            "Residence": "ARC", "PAYOR?": "SPONSOR", "SPONSOR NAME": "PETRONAS", "Remarks": "B3"}))
        self.assertEqual((rows[0]["Paid by"], rows[0]["Payor name"]), ("Sponsor", "PETRONAS (B3)"))
        self.assertNotIn("remarks", status[2][5] or "")

    def test_other_remarks_still_need_reading(self):
        _, _, rows, status, _ = self.convert(resident(**{
            "Residence": "ARC", "PAYOR?": "SPONSOR", "SPONSOR NAME": "MARA", "Remarks": "nak move ke Room 1"}))
        self.assertEqual(rows[0]["Payor name"], "MARA")
        self.assertIn("remarks", status[2][5])


class ApplicationMatchTests(unittest.TestCase):
    def app(self, row, stamp, email, phone, name, document_id="", student_id="", **values):
        base = {
            "how did you hear about us": "University",
            "marital status": "Single",
            "race": "",
            "race 2": "Indian",
            "religion": "Hindu",
            "residential address": "1 Test Street",
            "postal code": 12345,
            "state/ region/ province": "Test State",
            "country": "India",
            "do you have any medical conditions we should be aware of?": "No",
            'if you answered "yes" to the question above, please specify': "ignored when No",
            "please upload a colour passport size photograph of yourself": "https://example.test/photo",
            "please upload a colour copy of your nric (for malaysians only) or passport information page": "https://example.test/id",
            "please upload your university offer letter": "https://example.test/offer",
        }
        base.update(values)
        return {
            "_row": row,
            "_timestamp": stamp,
            "_document_id": cb.normalized_identifier(document_id),
            "_student_id": cb.normalized_identifier(student_id),
            "_email": cb.normalized_email(email),
            "_phone": cb.normalized_phone(phone),
            "_name": cb.normalized_name(name),
            "_values": base,
        }

    def resident(self, **changes):
        out = {
            "Email": "one@example.test",
            "Country Code": "+60",
            "Mobile Number": "123456789",
            "Resident Name": "Tést Person-One",
        }
        out.update(changes)
        return out

    def test_name_matching_ignores_accents_case_and_punctuation(self):
        self.assertEqual(cb.normalized_name("Tést Person-One"), cb.normalized_name("test person one"))

    def test_identifier_matching_ignores_case_and_punctuation(self):
        self.assertEqual(cb.normalized_identifier("a-123 45"), "A12345")
        self.assertEqual(cb.normalized_identifier("N/A"), "")

    def test_nric_or_passport_can_match_by_itself(self):
        rows = [self.app(10, datetime(2026, 1, 1), "other@example.test", "+999999", "Different Name",
                         document_id="A-123 45")]
        resident = self.resident(Email="", **{"Country Code": "", "Mobile Number": "", "Resident Name": "",
                                              "Passport number": "a12345"})
        matched, method, finding = cb.match_application(resident, cb.application_indexes(rows))
        self.assertEqual(matched["_row"], 10)
        self.assertEqual(method, "NRIC/passport")
        self.assertIsNone(finding)

    def test_university_student_id_can_match_when_the_source_has_that_field(self):
        rows = [self.app(10, datetime(2026, 1, 1), "other@example.test", "+999999", "Different Name",
                         student_id="U-2504322")]
        resident = self.resident(Email="", **{"Country Code": "", "Mobile Number": "", "Resident Name": "",
                                              "University Student ID": "u2504322"})
        matched, method, finding = cb.match_application(resident, cb.application_indexes(rows))
        self.assertEqual(matched["_row"], 10)
        self.assertEqual(method, "university Student ID")
        self.assertIsNone(finding)

    def test_quickbooks_id_is_not_treated_as_university_student_id(self):
        rows = [self.app(10, datetime(2026, 1, 1), "other@example.test", "+999999", "Different Name",
                         student_id="00305")]
        resident = self.resident(Email="", **{"Country Code": "", "Mobile Number": "", "Resident Name": "",
                                              "QuickBooks ID": "00305"})
        matched, method, finding = cb.match_application(resident, cb.application_indexes(rows))
        self.assertIsNone(matched)
        self.assertEqual(method, "No match")
        self.assertEqual(finding[0], cb.WARNING)  # no form is not a stopper: the Application Form Link is sent after import

    def test_one_identity_field_is_enough_when_it_is_unambiguous(self):
        rows = [self.app(10, datetime(2026, 1, 1), "one@example.test", "+999999", "Different Name")]
        matched, method, finding = cb.match_application(self.resident(), cb.application_indexes(rows))
        self.assertEqual(matched["_row"], 10)
        self.assertEqual(method, "email")
        self.assertIsNone(finding)

    def test_latest_repeat_submission_wins_quietly(self):
        rows = [
            self.app(10, datetime(2026, 1, 1), "one@example.test", "+60123456789", "Test Person One"),
            self.app(11, datetime(2026, 2, 1), "one@example.test", "+60123456789", "Test Person One"),
        ]
        matched, method, finding = cb.match_application(self.resident(), cb.application_indexes(rows))
        self.assertEqual(matched["_row"], 11)
        self.assertIn("newest of rows 10, 11", method)
        self.assertIsNone(finding)

    def test_two_agreeing_fields_beat_one_conflicting_field_quietly(self):
        rows = [
            self.app(10, datetime(2026, 1, 1), "one@example.test", "+999999", "Test Person One"),
            self.app(11, datetime(2026, 1, 1), "other@example.test", "+60123456789", "Someone Else"),
        ]
        matched, method, finding = cb.match_application(self.resident(), cb.application_indexes(rows))
        self.assertEqual(matched["_row"], 10)
        self.assertIn("also weakly matched row(s) 11", method)
        self.assertIsNone(finding)

    def test_one_against_one_is_ambiguous_and_not_guessed(self):
        rows = [
            self.app(10, datetime(2026, 1, 1), "one@example.test", "+999999", "Different Name"),
            self.app(11, datetime(2026, 1, 1), "other@example.test", "+60123456789", "Someone Else"),
        ]
        matched, _, finding = cb.match_application(self.resident(), cb.application_indexes(rows))
        self.assertIsNone(matched)
        self.assertEqual(finding[0], cb.BLOCKED)

    def test_only_requested_application_fields_and_documents_are_copied(self):
        row = self.app(10, datetime(2026, 1, 1), "one@example.test", "+60123456789", "Test Person One")
        out, findings = {}, []
        cb.enrich_from_application(out, row, Path("."), lambda *finding: findings.append(finding))
        self.assertEqual(out["Race"], "Indian")
        self.assertEqual(out["Religion"], "Hinduism")
        self.assertEqual(out["Postal Code"], "12345")
        self.assertEqual(out["Medical Condition - Remarks"], "")
        self.assertEqual(out["Photograph"], "https://example.test/photo")
        self.assertNotIn("Intake", out)
        self.assertEqual(findings, [])

    def test_corrections_file_can_assign_an_unmatched_application_row(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "corrections.json"
            path.write_text('{"applicationRows": {"00305": 123}}')
            self.assertEqual(cb.load_corrections(path)["applicationRows"], {"00305": 123})


class InvoiceNumberTests(ConverterTests):
    def test_original_invoice_number_comes_from_the_copy_name(self):
        _, _, rows, _, _ = self.convert(resident(**{"Initial Payment Invoice Copy": ("INV-IP_00894 .pdf", "https://drive.google.com/file/d/x")}))
        self.assertEqual(rows[0]["Invoice No."], "INV:IP_00894")


class ApplicationDropdownTests(unittest.TestCase):
    """Form answers become the system's own dropdown values - never new text (Dani, 13 Oct 2026)."""

    def test_marital_status_maps_or_stays_blank(self):
        self.assertEqual(cb.normal_marital("single")[0], "Single")
        self.assertEqual(cb.normal_marital("Bujang")[0], "Single")
        self.assertEqual(cb.normal_marital("widower")[0], "Widowed")
        value, note = cb.normal_marital("It's complicated")
        self.assertEqual(value, "")
        self.assertIn("left blank", note)
        self.assertEqual(cb.normal_marital(""), ("", ""))

    def test_referral_source_fits_the_application_form_list(self):
        self.assertEqual(cb.normal_referral("From my university")[0], "University")
        self.assertEqual(cb.normal_referral("Google")[0], "Google Search")
        self.assertEqual(cb.normal_referral("mmu email")[0], "University")
        self.assertEqual(cb.normal_referral("Website")[0], "Google Search")
        self.assertEqual(cb.normal_referral("Qstudy agent")[0], "Agent")
        self.assertEqual(cb.normal_referral("My senior told me")[0], "Friend")
        self.assertEqual(cb.normal_referral("Instagram")[0], "Social Media")
        self.assertEqual(cb.normal_referral("Agent")[0], "Agent")
        value, note = cb.normal_referral("Walked past the building")
        self.assertEqual(value, "Other")
        self.assertIn("grouped as Other", note)

    def test_medical_condition_is_yes_or_no_only(self):
        self.assertEqual(cb.normal_medical("Yes")[0], "yes")
        self.assertEqual(cb.normal_medical("No")[0], "no")
        self.assertEqual(cb.normal_medical("none")[0], "no")
        value, note = cb.normal_medical("Asthma")
        self.assertEqual(value, "")
        self.assertIn("not Yes or No", note)

    def test_every_answer_is_on_the_system_list(self):
        allowed = cb.allowed_values(cb.DEFAULT_LABELS)
        for fn, field, samples in (
            (cb.normal_marital, "Marital status", ["single", "married", "divorced", "widow"]),
            (cb.normal_referral, "Referral source", ["uni", "google", "friend", "agent", "tiktok", "other", "???"]),
            (cb.normal_medical, "Medical condition", ["yes", "no"]),
            (cb.normal_race, "Race", ["malay", "chinese", "bamar", "something else"]),
            (cb.normal_religion, "Religion", ["muslim", "kristian", "none", "something else"]),
        ):
            for sample in samples:
                self.assertIn(fn(sample)[0], allowed[field], f"{field}: {sample}")


class LegacyAgreementNoTests(ConverterTests):
    """An existing resident keeps their TA's own number (Rina, 13 Oct 2026)."""

    def test_format(self):
        self.assertEqual(cb.legacy_agreement_no(date(2025, 9, 1), "A-18-05", "00724"), "TA2025_A-18-05_00724")

    def test_blank_when_a_part_is_missing(self):
        self.assertEqual(cb.legacy_agreement_no(None, "A-18-05", "00724"), "")
        self.assertEqual(cb.legacy_agreement_no(date(2025, 9, 1), "", "00724"), "")

    def test_sheet3_is_filled(self):
        _, _, rows, _, _ = self.convert(resident())
        no = rows[0]["Agreement No."]
        self.assertRegex(no, r"^TA\d{4}_[A-Z0-9-]+_\d{5}$")
        self.assertTrue(no.endswith("_" + rows[0]["QuickBooks ID"]))
        self.assertIn("_" + rows[0]["Unit no. "] + "_", no)


class LaterBatchTests(unittest.TestCase):
    """Batch 2 onward: same converter, the master's own words accepted (Dani, 14 Oct 2026)."""

    def test_nationality_words_are_the_systems_codes(self):
        self.assertEqual(cb.NATIONALITY_WORDS["MOROCCAN"], "MAR")
        self.assertEqual(cb.NATIONALITY_WORDS["MALAYSIAN"], "MYS")
        allowed = cb.allowed_values(cb.DEFAULT_LABELS)["Nationality"]
        for word in ("MOROCCAN", "INDIAN", "INDONESIAN"):
            self.assertIn(cb.NATIONALITY_WORDS[word], allowed)

    def test_more_race_and_religion_words(self):
        self.assertEqual(cb.normal_race("caucasian")[0], "European")
        self.assertEqual(cb.normal_religion("Atheist")[0], "No Religion / None")

    def test_semi_annually_is_a_frequency(self):
        self.assertEqual(cb.FREQUENCY["SEMI-ANNUALLY"], "semiannual")


class MasterQuirkTests(unittest.TestCase):
    """What the master really holds (Dani, 14 Oct 2026)."""

    def test_apostrophe_kept_by_excel_is_not_part_of_the_number(self):
        self.assertEqual(cb.split_phone("+962 791234567'")[:2], ("+962", "791234567"))
        self.assertEqual(cb.split_phone("'0123456789")[:2], ("+60", "123456789"))

    def test_country_code_without_plus_from_the_nationality(self):
        code, mobile, finding = cb.split_phone("263771234567", "ZWE")
        self.assertEqual((code, mobile), ("+263", "771234567"))
        self.assertEqual(finding[0], cb.WARNING)

    def test_unknown_number_still_needs_review(self):
        self.assertEqual(cb.split_phone("12345", "ZWE")[2][0], cb.REVIEW)


class AdminNextRentTests(ConverterTests):
    """Brachtia admin's next payment wins over the worked-out one (Dani, 15 Oct 2026)."""

    def test_corrections_next_rent_is_written_into_sheet3(self):
        fix = self.tmp / "fix.json"
        qb = resident()["StudentID"]
        fix.write_text(json.dumps({"nextRent": {str(qb).zfill(5): {"due": "2026-12-05", "amount": 999.5}}}))
        _, _, rows, _, _ = Fixture([resident()], self.tmp).run(fix)
        self.assertEqual(rows[0]["Next payment amount"], 999.5)
        self.assertEqual(str(rows[0]["Next payment due"])[:10], "2026-12-05")


class WholeMonthFlagTests(ConverterTests):
    """What the whole-month rule cannot read on its own is flagged (Dani, 15 Oct 2026)."""

    # both signs mostly mean a renewal: review until the team gives the next payment (16 Oct 2026)
    def test_start_on_the_1st_with_a_pro_rated_amount(self):
        _, _, _, status, _ = self.convert(resident(**{"Tenancy Start": datetime(2026, 9, 1)}))
        self.assertEqual(status[2][3], "Needs review")

    def test_advance_not_whole_months(self):
        _, _, _, status, _ = self.convert(resident(**{"Advance Rental": 1000}))
        self.assertEqual(status[2][3], "Needs review")

    def test_team_next_payment_turns_review_into_warning(self):
        fix = self.tmp / "fix.json"
        qb = str(resident()["StudentID"]).zfill(5)
        fix.write_text(json.dumps({"nextRent": {qb: {"due": "2026-12-05", "amount": 1000}}}))
        _, _, _, status, _ = Fixture([resident(**{"Advance Rental": 1000})], self.tmp).run(fix)
        self.assertEqual(status[2][3], "Warning")
