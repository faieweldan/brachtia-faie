"""
A second pair of eyes on the rows Brachtia marked "Checked" - for people, not the upload.

  python3 scripts/review-checked-rows.py "<master list.xlsx>"
  python3 scripts/review-checked-rows.py "<master list.xlsx>" --today 2026-09-17

The master list's DB MIGRATION Status (column AH) runs Proceed -> Checked -> Cleaned.
Before a row is marked Cleaned, someone should confirm the work behind "Checked".
This script lists what looks wrong so that person knows where to look. It decides
nothing and removes nothing: every Checked row is kept, and the doubtful cells are
coloured.

Only Checked rows are reviewed, but they are compared against EVERY row in the
sheet - a Checked row can clash with a Proceed row on the same bed.

Writes "<master list> - checked review.xlsx" next to the master list, with:

  Read me first   the colours, how many cells and rows, who resolves each
  Student Data    the Checked rows, problem cells coloured, problems listed at the end
                  (monthly payment columns left out)
  Issue list      one line per coloured cell: Excel row, field, problem, value

Read only: the master list is never written, and no database is touched.
The terminal prints counts only - no names, no ids.

This is not the upload's dry run (scripts/check-master-list.ts). That one asks
"what would the upload do?"; this one asks "is the sheet telling the truth?".
"""

import datetime
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from statistics import median

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter

# ---------------- arguments ----------------

args = sys.argv[1:]
today = datetime.date.today()
if "--today" in args:
    i = args.index("--today")
    today = datetime.date.fromisoformat(args[i + 1])
    del args[i : i + 2]
if not args:
    print('Usage: python3 scripts/review-checked-rows.py "<master list.xlsx>" [--today YYYY-MM-DD]')
    sys.exit(1)
src = Path(args[0])
out = Path(args[1]) if len(args) > 1 else src.with_name(f"{src.stem} - checked review.xlsx")

SHEET = "Student Data"
REVIEW_STAGE = "checked"

# ---------------- the checks ----------------

# level -> (fill, meaning, who resolves). Order = how much it blocks.
LEVELS = {
    "clash": ("FFC7CE", "Two people on one bed, or one student in two beds", "Ops - only they know who actually lives there"),
    "dates": ("FFD9B3", "Tenancy date missing or broken", "Ops - check the tenancy agreement"),
    "status": ("FFF2CC", "Status doesn't match the dates", "Ops - check move-in / move-out"),
    "person": ("DDEBF7", "Student ID or details missing", "Ops - confirm with the resident"),
    "money": ("E2EFDA", "Rent, deposit or frequency missing", "Accounting - check the payment record"),
}
BLOCKING = "Red, orange, yellow and blue block the migration. Green does not."

# key -> (level, cells to colour, short name, one-line meaning)
CHECKS = {
    "bed-overlap": ("clash", ["Bed"], "Bed double-booked", "Two tenancies on one bed, same dates"),
    "room-two-ways": ("clash", ["Bed"], "Single and twin together", "Room let both ways, same dates"),
    "unit-and-rooms": ("clash", ["Bed"], "Whole unit and rooms together", "Unit let whole and by room, same dates"),
    "id-two-places": ("clash", ["StudentID"], "Student in two beds", "Same StudentID, same dates, two rows"),
    "id-two-names": ("clash", ["StudentName"], "One ID, two names", "Same StudentID, different name"),
    "unitid-mismatch": ("clash", ["UnitID"], "UnitID mismatch", "One UnitID, two unit numbers"),
    "active-ended": ("status", ["Status"], "Active but ended", "Should be Inactive?"),
    "active-future": ("status", ["Status"], "Active but not started", "Should be Booked?"),
    "inactive-current": ("status", ["Status"], "Inactive but still in tenancy", "Left early, or still Active?"),
    "booked-started": ("status", ["Status"], "Booked but started", "Moved in, or cancelled?"),
    "vacant-has-student": ("status", ["Status"], "Vacant with a student", "Vacant row has ID, dates or rent"),
    "no-status": ("status", ["Status"], "No status", "Status is empty"),
    "date-broken": ("dates", ["Tenancy Start", "Tenancy End"], "Bad date", "Can't be read"),
    "date-missing": ("dates", ["Tenancy Start", "Tenancy End"], "Missing date", "No start or end"),
    "end-first": ("dates", ["Tenancy End"], "End before start", "End date is earlier"),
    "id-missing": ("person", ["StudentID"], "No StudentID", "Student has no ID"),
    "id-format": ("person", ["StudentID"], "ID format", "Not five digits, or several IDs"),
    "info-missing": ("person", [], "Info missing", "Empty student details"),
    "gender-odd": ("person", ["Gender"], "Gender not M/F", "Written another way"),
    "no-rent": ("money", ["Monthly Rent"], "No rent", "Monthly rent empty"),
    "no-frequency": ("money", ["Payment Frequency"], "No frequency", "Payment frequency empty"),
    "no-deposit": ("money", ["Security Deposit"], "No deposit", "Security deposit empty (write 0 if none)"),
    "paid-over-total": ("money", ["Total Initial Payment Made"], "Paid over total", "Made is more than total"),
}

# ---------------- reading ----------------

wb_values = openpyxl.load_workbook(src, data_only=True)
if SHEET not in wb_values.sheetnames:
    print(f'No "{SHEET}" sheet in {src.name}')
    sys.exit(1)
sheet = wb_values[SHEET]


def norm(v):
    return str(v if v is not None else "").strip()


header_row = next(
    (row[0].row for row in sheet.iter_rows(max_row=30) if any(norm(c.value).lower() == "studentid" for c in row)),
    None,
)
if header_row is None:
    print("No StudentID header - this is not the master list.")
    sys.exit(1)

headers = [c.value for c in sheet[header_row]]
width = len(headers)


def header_text(h):
    return h.strftime("%b %Y") if isinstance(h, (datetime.date, datetime.datetime)) else norm(h)


# the first column with a name is the live one; a repeated name is the "second" copy
first_col, second_col = {}, {}
for i, h in enumerate(headers):
    name = header_text(h).lower()
    if not name:
        continue
    if name not in first_col:
        first_col[name] = i
    elif name not in second_col:
        second_col[name] = i


def col(name, second=False):
    table = second_col if second else first_col
    return table.get(name.lower())


REQUIRED = ["unitid", "unit", "room", "bed", "studentid", "studentname", "tenancy start", "tenancy end", "status",
            "db migration status"]
missing = [n for n in REQUIRED if n not in first_col]
if missing:
    print("Master list is missing columns: " + ", ".join(missing))
    sys.exit(1)


class Row:
    def __init__(self, sheet_row, cells):
        self.sheet_row = sheet_row
        self.cells = cells

    def get(self, name, second=False):
        i = col(name, second)
        return None if i is None or i >= len(self.cells) else self.cells[i]

    def text(self, name, second=False):
        return norm(self.get(name, second))


rows = []
for cells in sheet.iter_rows(min_row=header_row + 1, max_col=width):
    values = [c.value for c in cells]
    if any(norm(v) for v in values):
        rows.append(Row(cells[0].row, values))

ERROR = re.compile(r"^#(VALUE!|REF!|N/A|DIV/0!|NAME\?|NUM!|NULL!)$")


def read_date(v):
    """(date or None, broken?)"""
    if v is None or norm(v) in ("", "-"):
        return None, False
    if isinstance(v, datetime.datetime):
        v = v.date()
    if isinstance(v, datetime.date):
        return (v, False) if 2000 <= v.year <= 2099 else (None, True)
    return None, True  # text, an Excel error, or a bare number


def money(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^0-9.\-]", "", norm(v))
    try:
        return float(s) if s else None
    except ValueError:
        return None


for r in rows:
    r.stage = r.text("db migration status").lower()
    r.status = r.text("status").lower()
    r.ids = norm(r.get("studentid")).split()
    r.name = re.sub(r"\s+", " ", r.text("studentname")).lower()
    r.unit, r.room, r.bed = r.text("unit").upper(), r.text("room").upper(), r.text("bed").lower()
    r.start, r.start_broken = read_date(r.get("tenancy start"))
    r.end, r.end_broken = read_date(r.get("tenancy end"))
    r.is_vacant = r.status == "vacant"
    r.has_student = bool(r.ids) or (r.status in ("active", "inactive", "booked"))
    r.shape = "unit" if r.bed == "unit" or r.room == "UNIT" else ("single" if r.bed == "single" else "twin")


def overlaps(a, b):
    return a.start and a.end and b.start and b.end and a.start <= b.end and b.start <= a.end


def where(r):
    return f"{r.unit} / {r.room} / {r.text('bed')}"


def ref(r):
    stage = r.text("db migration status") or "no stage"
    return f"row {r.sheet_row} ({stage})"


# ---------------- finding ----------------

findings = []  # (row, check key, message, cells to colour)


def flag(r, key, message, fields=None):
    findings.append((r, key, message, fields or CHECKS[key][1]))


reviewed = [r for r in rows if r.stage == REVIEW_STAGE]
tenancies = [r for r in rows if r.has_student and not r.is_vacant]

by_bed, by_room, by_unit, by_id, by_name = (defaultdict(list) for _ in range(5))
for r in tenancies:
    by_bed[(r.unit, r.room, r.bed)].append(r)
    by_room[(r.unit, r.room)].append(r)
    by_unit[r.unit].append(r)
    for sid in r.ids:
        by_id[sid].append(r)
    if r.name and r.name != "vacant":
        by_name[r.name].append(r)

unit_ids = defaultdict(set)
for r in rows:
    if r.unit:
        unit_ids[r.text("unitid")].add(r.unit)


def d(x):
    return f"{x:%d %b %y}"


for r in reviewed:
    occupied = r.has_student and not r.is_vacant

    # --- clash ---
    if occupied:
        for other in by_bed[(r.unit, r.room, r.bed)]:
            if other is not r and overlaps(r, other):
                hint = ""
                # the usual cause: both twins typed as Twin 1 while Twin 2 sits Vacant
                if r.shape == "twin":
                    twin = "Twin 2" if r.bed == "twin 1" else "Twin 1"
                    if any(v.is_vacant and (v.unit, v.room, v.bed) == (r.unit, r.room, twin.lower()) for v in rows):
                        hint = f" - {twin} is Vacant, one should be {twin}?"
                flag(r, "bed-overlap", f"Same bed as row {other.sheet_row}, dates overlap{hint}")
        if r.shape != "unit":
            for other in by_room[(r.unit, r.room)]:
                if other is not r and other.shape in ("single", "twin") and other.shape != r.shape and overlaps(r, other):
                    flag(r, "room-two-ways", f"Room let as {other.shape} in row {other.sheet_row}, dates overlap")
        for other in by_unit[r.unit]:
            if other is not r and (r.shape == "unit") != (other.shape == "unit") and overlaps(r, other):
                how = "whole" if other.shape == "unit" else "by room"
                flag(r, "unit-and-rooms", f"Unit let {how} in row {other.sheet_row}, dates overlap")
        for sid in r.ids:
            for other in by_id[sid]:
                if other is r:
                    continue
                if overlaps(r, other):
                    flag(r, "id-two-places", f"Same StudentID in row {other.sheet_row}, dates overlap")
                if len(r.ids) == 1 and len(other.ids) == 1 and r.name and other.name and r.name != other.name:
                    flag(r, "id-two-names", f"Same StudentID in row {other.sheet_row}, different name")
    uid = r.text("unitid")
    if uid and len(unit_ids[uid]) > 1:
        flag(r, "unitid-mismatch", f"UnitID also used for {', '.join(sorted(unit_ids[uid] - {r.unit}))}")

    # --- status ---
    if not r.status:
        flag(r, "no-status", "Status empty")
    elif r.is_vacant:
        has = [n for n, v in (("StudentID", r.ids), ("dates", r.start or r.end), ("rent", money(r.get("monthly rent")))) if v]
        if has:
            flag(r, "vacant-has-student", f"Vacant, but has {' and '.join(has)}")
    elif r.status == "active":
        if r.end and r.end < today:
            flag(r, "active-ended", f"Active, but ended {d(r.end)}")
        elif r.start and r.start > today:
            flag(r, "active-future", f"Active, but starts {d(r.start)}")
    elif r.status == "inactive":
        if r.start and r.end and r.start <= today <= r.end:
            flag(r, "inactive-current", f"Inactive, but tenancy runs to {d(r.end)}")
    elif r.status == "booked":
        if r.start and r.start <= today:
            flag(r, "booked-started", f"Booked, but started {d(r.start)}")

    # --- dates ---
    for label, broken, cell in (("Tenancy Start", r.start_broken, r.get("tenancy start")),
                                ("Tenancy End", r.end_broken, r.get("tenancy end"))):
        if broken:
            shown = cell.date() if isinstance(cell, datetime.datetime) else cell
            flag(r, "date-broken", f'{label} can\'t be read ("{shown}")', [label])
    if occupied:
        for label, broken, value in (("Tenancy Start", r.start_broken, r.start), ("Tenancy End", r.end_broken, r.end)):
            if not broken and not value:
                flag(r, "date-missing", f"{label} empty", [label])
    if r.start and r.end and r.end < r.start:
        flag(r, "end-first", "Tenancy ends before it starts")

    # --- person ---
    if occupied:
        if not norm(r.get("studentid")):
            flag(r, "id-missing", "StudentID empty")
        elif len(r.ids) > 1:
            flag(r, "id-format", f"{len(r.ids)} StudentIDs in one cell")
        elif not re.fullmatch(r"\d{5}", str(r.get("studentid"))):
            spaces = re.fullmatch(r"\d{5}", norm(r.get("studentid")))
            flag(r, "id-format", "StudentID has extra spaces" if spaces else "StudentID not five digits")
        for label, name in (("Gender", "gender"), ("University", "university"), ("Nationality", "nationality"),
                            ("Sponsor", "sponsor"), ("ID number", "idnumber"), ("Mobile", "mobilenumber"),
                            ("Email", "email")):
            if col(name) is not None and not r.text(name):
                flag(r, "info-missing", f"{label} empty", [name])
        g = r.text("gender")
        if g and g not in ("M", "F"):
            flag(r, "gender-odd", f'Gender written as "{g}", not M/F')

    # --- money ---
    if occupied:
        if not money(r.get("monthly rent")):
            flag(r, "no-rent", "Monthly rent empty")
        if col("payment frequency") is not None and not r.text("payment frequency"):
            flag(r, "no-frequency", "Payment frequency empty")
        if col("security deposit") is not None and money(r.get("security deposit")) is None:
            flag(r, "no-deposit", "Security deposit empty")
        total, made = money(r.get("total initial payment")), money(r.get("total initial payment made"))
        if total is not None and made is not None and made > total + 0.01:
            flag(r, "paid-over-total", f"Initial payment made ({made:,.0f}) is more than the total ({total:,.0f})")

# ---------------- drawing ----------------


def solid(hex_colour):
    return PatternFill("solid", fgColor="FF" + hex_colour)


BOLD = Font(bold=True)
order = list(CHECKS)
level_rank = {lvl: i for i, lvl in enumerate(LEVELS)}
rows_hit = defaultdict(set)
for r, k, _, _ in findings:
    rows_hit[k].add(r.sheet_row)

per_row = defaultdict(list)
for r, k, m, f in findings:
    per_row[r.sheet_row].append((k, m, f))

HEAD_FILL, HEAD_FONT = solid("44546A"), Font(bold=True, color="FFFFFFFF")


def head(ws):
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT


book = openpyxl.Workbook()
readme = book.active
readme.title = "Read me first"
grid = book.create_sheet("Student Data")
issues = book.create_sheet("Issue list")

# the monthly payment columns (headed by a date) are left out - see the note
shown = [i for i, h in enumerate(headers) if not isinstance(h, (datetime.date, datetime.datetime))]
months_hidden = len(headers) - len(shown)
out_col = {i: 3 + n for n, i in enumerate(shown)}

# Student Data
grid.append(["Flag", "Excel row"] + [header_text(headers[i]) for i in shown] + ["Issues on this row"])
head(grid)
cells_hit, level_rows, issue_lines = Counter(), defaultdict(set), []

for r in reviewed:
    doubts = sorted(per_row.get(r.sheet_row, []), key=lambda x: (level_rank[CHECKS[x[0]][0]], order.index(x[0])))
    # one message once, even when two rows point at each other more than once
    messages = list(dict.fromkeys(m for _, m, _ in doubts))
    grid.append(["!" if doubts else "", r.sheet_row] + [r.cells[i] for i in shown] + ["; ".join(messages)])
    out_row = grid.max_row
    for i in shown:
        if isinstance(r.cells[i], (datetime.date, datetime.datetime)):
            grid.cell(out_row, out_col[i]).number_format = "d-mmm-yy"

    cell_problems = defaultdict(list)
    for k, message, fields in doubts:
        lvl = CHECKS[k][0]
        level_rows[lvl].add(r.sheet_row)
        for n in fields:
            i = col(n)
            if i is not None and i in out_col and message not in (m for _, m in cell_problems[i]):
                cell_problems[i].append((lvl, message))
    for i, problems in cell_problems.items():
        worst = min(problems, key=lambda n: level_rank[n[0]])[0]
        cells_hit[worst] += 1
        grid.cell(out_row, out_col[i]).fill = solid(LEVELS[worst][0])
        value = r.cells[i]
        value = value.strftime("%d %b %Y") if isinstance(value, (datetime.date, datetime.datetime)) else norm(value)
        issue_lines.append((level_rank[worst], r.sheet_row, out_col[i], header_text(headers[i]),
                            "; ".join(m for _, m in problems), value, LEVELS[worst][0]))

grid.freeze_panes = "H2"
grid.auto_filter.ref = f"A1:{get_column_letter(grid.max_column)}{grid.max_row}"
for letter, w in (("A", 6), ("B", 9)):
    grid.column_dimensions[letter].width = w
grid.column_dimensions[get_column_letter(grid.max_column)].width = 60

# Read me first
readme.append(["Brachtia Student Data - Checked rows review"])
readme["A1"].font = Font(bold=True, size=14)
readme.append([])
readme.append(["Only the problem CELLS are coloured, so we can see exactly which field is wrong. "
               "Column A flags any row with at least one problem - filter on it."])
readme.append(["The last column lists every problem on that row. Fix in the master sheet, not here."])
readme.append([])
readme.append([f"{len(reviewed)} Checked rows, {len({f[0].sheet_row for f in findings})} with a problem. "
               f"As of {today:%d %b %Y}."])
readme.append([])
readme.append(["Colour", "Meaning", "Cells", "Rows", "Who resolves"])
for c in readme[readme.max_row]:
    c.font = BOLD
for lvl, (fill, meaning, who) in LEVELS.items():
    if level_rows[lvl]:
        readme.append(["", meaning, cells_hit[lvl], len(level_rows[lvl]), who])
        readme.cell(readme.max_row, 1).fill = solid(fill)
readme.append([])
readme.append([BLOCKING])
readme.cell(readme.max_row, 1).font = BOLD
readme.append([])
readme.append([f"Note: the {months_hidden} monthly payment columns are not shown here."])
for letter, w in zip("ABCDE", (10, 46, 8, 8, 50)):
    readme.column_dimensions[letter].width = w

# Issue list: one line per coloured cell
issues.append(["Excel row", "Field", "Problem", "Value"])
head(issues)
for _, row_no, _, field, problem, value, fill in sorted(issue_lines):
    issues.append([row_no, field, problem, value])
    issues.cell(issues.max_row, 2).fill = solid(fill)
issues.freeze_panes = "A2"
issues.auto_filter.ref = f"A1:D{issues.max_row}"
for letter, w in zip("ABCD", (11, 18, 48, 60)):
    issues.column_dimensions[letter].width = w

book.save(out)

# ---------------- terminal: counts only ----------------

print(f"\n{len(reviewed)} Checked rows, {len({f[0].sheet_row for f in findings})} with issues\n")
for key in order:
    if rows_hit[key]:
        print(f"  {CHECKS[key][2]:<32} {len(rows_hit[key]):>4}")
print(f"\n{out}")
