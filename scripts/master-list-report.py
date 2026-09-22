"""
Draw the upload check as a coloured copy of the master list.

Called by scripts/check-master-list.ts, which decides what is wrong with each
row using the same rules as Bulk upload. This file decides nothing - it only
draws what it is given:

  Read me first   the colours, what each means, how many, and who resolves it
  Checked rows    the student rows, with only the problem cells coloured
  Issue list      one line per problem, to filter and sort

  python3 scripts/master-list-report.py <master.xlsx> <findings.json> <out.xlsx>

Why Python: the xlsx library the app uses cannot write cell colours. openpyxl
can, and it is already on the Mac - so the app gains no dependency for a report.
"""

import datetime
import json
import sys

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

src, findings_path, out = sys.argv[1:4]
with open(findings_path, encoding="utf-8") as fh:
    data = json.load(fh)

groups = data["groups"]
findings = data["findings"]
rank = {g["key"]: i for i, g in enumerate(groups)}
group_of = {g["key"]: g for g in groups}


def solid(hex_colour):
    return PatternFill("solid", fgColor="FF" + hex_colour)


def norm(value):
    return str(value or "").strip().lower()


HEAD_FILL = solid("44546A")
HEAD_FONT = Font(bold=True, color="FFFFFFFF")
MUTED = Font(color="FF5B6760")

# ---------------- the source sheet ----------------

master = openpyxl.load_workbook(src)
sheet = master[data["sheet"]]

header_row = next(
    row[0].row
    for row in sheet.iter_rows(max_row=30)
    if any(norm(c.value) == "studentid" for c in row)
)
headers = [c.value for c in sheet[header_row]]

# the monthly payment columns start at the first date header. Like the cleanup
# file, they are left out: the upload does not read them
last = next(
    (i for i, h in enumerate(headers) if isinstance(h, (datetime.date, datetime.datetime))),
    len(headers),
)
while last and norm(headers[last - 1]) == "":
    last -= 1
headers = headers[:last]

# a repeated header name (Status, Tenancy Start...) means its FIRST column - the
# one the upload reads
first_col = {}
for i, h in enumerate(headers):
    first_col.setdefault(norm(h), i)

# which cells to colour, and the worst colour on each row
cell_marks = {}
row_worst = {}
for f in findings:
    r = f["sheetRow"]
    for field in f["fields"]:
        c = first_col.get(norm(field))
        if c is not None:
            cell_marks.setdefault((r, c), []).append(f)
    worst = row_worst.get(r)
    if worst is None or rank[f["group"]] < rank[worst]:
        row_worst[r] = f["group"]

book = openpyxl.Workbook()

# ---------------- Read me first ----------------

readme = book.active
readme.title = "Read me first"
s = data["summary"]
readme["A1"] = "Master list - upload check"
readme["A1"].font = Font(bold=True, size=14)
readme["A2"] = data["subtitle"]
readme["A2"].font = MUTED
readme["A3"] = (
    "Nothing was uploaded. Only the problem cells are coloured, so you can see "
    "exactly which field is wrong. Column A of Checked rows shows the worst colour on each row."
)
readme["A5"] = (
    f'{s["rows"]} rows read  ·  {s["imported"]} would be imported  ·  '
    f'{s["placed"]} placed in a bed  ·  {s["emptied"]} beds emptied'
)
readme["A5"].font = Font(bold=True)

r = 7
for col, text in zip("ABCDE", ["Colour", "Meaning", "Rows", "Why", "Who resolves"]):
    readme[f"{col}{r}"] = text
    readme[f"{col}{r}"].font = Font(bold=True)

for g in groups:
    in_group = [f for f in findings if f["group"] == g["key"]]
    if not in_group:
        continue
    r += 1
    readme[f"A{r}"].fill = solid(g["fill"])
    readme[f"B{r}"] = g["meaning"]
    readme[f"B{r}"].font = Font(bold=True)
    readme[f"C{r}"] = len({f["sheetRow"] for f in in_group})
    readme[f"E{r}"] = g["who"]
    for p in g["patterns"]:
        count = sum(1 for f in in_group if f["pattern"] == p["key"])
        if not count:
            continue
        r += 1
        readme[f"B{r}"] = f'    {p["title"]}'
        readme[f"C{r}"] = count
        readme[f"D{r}"] = p["why"]
        readme[f"D{r}"].font = MUTED

r += 2
readme[f"A{r}"] = data["note"]
readme[f"A{r}"].font = Font(bold=True)

for col, width in {"A": 10, "B": 44, "C": 8, "D": 70, "E": 46}.items():
    readme.column_dimensions[col].width = width
for row in readme.iter_rows():
    for c in row:
        c.alignment = Alignment(vertical="top", wrap_text=c.column_letter in "DE")

# ---------------- Checked rows ----------------

rows_sheet = book.create_sheet("Checked rows")
rows_sheet.append(["Flag", "Excel row", *[str(h or "") for h in headers]])
for c in rows_sheet[1]:
    c.fill = HEAD_FILL
    c.font = HEAD_FONT

at = 1
for src_row in range(header_row + 1, sheet.max_row + 1):
    cells = [sheet.cell(src_row, i + 1) for i in range(len(headers))]
    if all(norm(c.value) == "" for c in cells):
        continue
    worst = row_worst.get(src_row)
    rows_sheet.append(
        [group_of[worst]["label"] if worst else "", src_row, *[c.value for c in cells]]
    )
    # counted, not read back: max_row scans the whole sheet each time it is asked
    at += 1
    if worst:
        rows_sheet.cell(at, 1).fill = solid(group_of[worst]["fill"])
    for i, source in enumerate(cells):
        target = rows_sheet.cell(at, i + 3)
        target.number_format = source.number_format
        marks = cell_marks.get((src_row, i))
        if marks:
            top = min(marks, key=lambda f: rank[f["group"]])
            target.fill = solid(group_of[top["group"]]["fill"])
            target.comment = Comment("\n".join(dict.fromkeys(f["problem"] for f in marks)), "Upload check")

name_col = first_col.get("studentname", 5) + 3
rows_sheet.freeze_panes = f"{get_column_letter(name_col + 1)}2"
rows_sheet.auto_filter.ref = rows_sheet.dimensions
rows_sheet.column_dimensions["A"].width = 14
rows_sheet.column_dimensions["B"].width = 9
for i, h in enumerate(headers):
    rows_sheet.column_dimensions[get_column_letter(i + 3)].width = (
        32 if norm(h) == "studentname" else max(10, min(24, len(str(h or "")) + 4))
    )

# ---------------- Issue list ----------------

issues = book.create_sheet("Issue list")
issues.append(["Excel row", "StudentID", "Field", "Problem", "Value"])
for c in issues[1]:
    c.fill = HEAD_FILL
    c.font = HEAD_FONT
ordered = sorted(findings, key=lambda f: (f["sheetRow"], rank[f["group"]]))
for issue_row, f in enumerate(ordered, start=2):
    field = f["fields"][0] if f["fields"] else ""
    c = first_col.get(norm(field))
    value = sheet.cell(f["sheetRow"], c + 1).value if c is not None else ""
    issues.append([f["sheetRow"], f["studentId"], field, f["problem"], value])
    issues.cell(issue_row, 1).fill = solid(group_of[f["group"]]["fill"])
issues.freeze_panes = "A2"
issues.auto_filter.ref = issues.dimensions
for col, width in {"A": 10, "B": 12, "C": 16, "D": 80, "E": 28}.items():
    issues.column_dimensions[col].width = width

book.save(out)
